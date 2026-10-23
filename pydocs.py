"""
Встроенный конвертер документов без pandoc и LibreOffice (телефон: встроенный Python).

Всё идёт через HTML: читаем DOCX (mammoth), Markdown, HTML, TXT, CSV/TSV, EPUB,
пишем HTML, Markdown (markdownify), TXT, DOCX (python-docx) и PDF (fpdf2).
Оформление попроще, чем у pandoc и LibreOffice, но заголовки, жирный, курсив, списки,
таблицы, ссылки и картинки сохраняются.
"""
from __future__ import annotations

import base64
import csv
import html as htmllib
import io
import os
import re
import zipfile
from pathlib import Path

READ = ("docx", "md", "markdown", "html", "htm", "txt", "csv", "tsv", "epub")
WRITE = ("html", "md", "txt", "docx")          # для конвертера; редактор ещё умеет PDF


class DocError(Exception):
    pass


def available() -> bool:
    try:
        import bs4  # noqa: F401
        import docx  # noqa: F401
        import mammoth  # noqa: F401
        import markdown  # noqa: F401
        return True
    except ImportError:
        return False


def _body(html: str) -> str:
    m = re.search(r"<body[^>]*>(.*)</body>", html, re.S | re.I)
    body = m.group(1) if m else html
    body = re.sub(r"<script\b.*?</script\s*>", "", body, flags=re.S | re.I)
    body = re.sub(r"<style\b.*?</style\s*>", "", body, flags=re.S | re.I)
    return re.sub(r"\son\w+\s*=\s*(\"[^\"]*\"|'[^']*'|[^\s>]+)", "", body, flags=re.I).strip()


# --------------------------------------------------------------------- чтение -> HTML

def read_html(src: Path) -> str:
    ext = src.suffix.lower().lstrip(".")
    raw = src.read_bytes()
    if ext == "docx":
        import mammoth
        with src.open("rb") as f:
            return mammoth.convert_to_html(f).value       # картинки сразу data: URI
    if ext in ("md", "markdown"):
        import markdown
        return markdown.markdown(raw.decode("utf-8", "replace"), extensions=["tables", "fenced_code", "sane_lists"])
    if ext in ("html", "htm"):
        return _body(raw.decode("utf-8", "replace"))
    if ext == "txt":
        text = raw.decode("utf-8", "replace")
        return "".join(f"<p>{htmllib.escape(l) or '<br>'}</p>" for l in text.splitlines())
    if ext in ("csv", "tsv"):
        rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig", "replace")), delimiter="\t" if ext == "tsv" else ","))
        cells = "".join("<tr>" + "".join(f"<td>{htmllib.escape(c)}</td>" for c in r) + "</tr>" for r in rows)
        return f"<table>{cells}</table>"
    if ext == "epub":
        parts = []
        with zipfile.ZipFile(src) as z:
            names = [n for n in z.namelist() if n.lower().endswith((".xhtml", ".html", ".htm"))]
            for n in sorted(names):
                parts.append(_body(z.read(n).decode("utf-8", "replace")))
        return "\n".join(parts)
    raise DocError(f"встроенный конвертер не читает .{ext}")


# --------------------------------------------------------------------- HTML -> форматы

def to_markdown(html: str) -> str:
    from markdownify import markdownify
    return markdownify(html, heading_style="ATX", bullets="-").strip() + "\n"


def to_text(html: str) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    for br in soup.find_all("br"):
        br.replace_with("\n")
    blocks = soup.find_all(["p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "tr", "pre", "blockquote"])
    if not blocks:
        return soup.get_text().strip() + "\n"
    lines = []
    for b in blocks:
        if b.find_parent(["li", "tr"]) and b.name not in ("li", "tr"):
            continue
        t = b.get_text(" " if b.name == "tr" else "", strip=b.name == "tr")
        lines.append(("• " if b.name == "li" else "") + t.strip())
    return "\n".join(lines).strip() + "\n"


def _style(tag) -> dict:
    out = {}
    for part in (tag.get("style") or "").split(";"):
        if ":" in part:
            k, v = part.split(":", 1)
            out[k.strip().lower()] = v.strip().lower()
    return out


def _data_image(src: str) -> io.BytesIO | None:
    m = re.match(r"data:image/[\w.+-]+;base64,(.*)", src or "", re.S)
    if not m:
        return None
    try:
        return io.BytesIO(base64.b64decode(m.group(1)))
    except ValueError:
        return None


def to_docx(html: str, out: Path, title: str = "") -> Path:
    """HTML -> DOCX: заголовки, абзацы, жирный/курсив/подчёркнутый/зачёркнутый, цвет, размер,
    выравнивание, списки, таблицы, ссылки (текстом), картинки data:."""
    from bs4 import BeautifulSoup, NavigableString, Tag
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Inches

    doc = Document()
    if title:
        doc.core_properties.title = title
    soup = BeautifulSoup(html, "html.parser")
    ALIGN = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}

    def runs(node, para, fmt):
        for ch in node.children:
            if isinstance(ch, NavigableString):
                text = str(ch)
                if not text.strip() and "\n" in text:
                    continue
                r = para.add_run(re.sub(r"\s+", " ", text))
                r.bold, r.italic, r.underline = fmt.get("b"), fmt.get("i"), fmt.get("u")
                r.font.strike = fmt.get("s")
                r.font.superscript, r.font.subscript = fmt.get("sup"), fmt.get("sub")
                if fmt.get("mono"):
                    r.font.name = "Courier New"
                if fmt.get("color"):
                    try:
                        r.font.color.rgb = RGBColor.from_string(fmt["color"])
                    except ValueError:
                        pass
                if fmt.get("size"):
                    r.font.size = Pt(fmt["size"])
                continue
            if not isinstance(ch, Tag):
                continue
            n, f = ch.name, dict(fmt)
            st = _style(ch)
            if n in ("b", "strong") or st.get("font-weight") in ("bold", "700", "600"):
                f["b"] = True
            if n in ("i", "em") or st.get("font-style") == "italic":
                f["i"] = True
            if n == "u" or "underline" in st.get("text-decoration", ""):
                f["u"] = True
            if n in ("s", "strike", "del") or "line-through" in st.get("text-decoration", ""):
                f["s"] = True
            if n in ("sup", "sub"):
                f[n] = True
            if n in ("code", "kbd", "tt"):
                f["mono"] = True
            if n == "a":
                f["u"] = True
                f["color"] = "1A73E8"
            m = re.match(r"#([0-9a-f]{6})", st.get("color", ""))
            if m:
                f["color"] = m.group(1).upper()
            m = re.match(r"([\d.]+)(pt|px)", st.get("font-size", ""))
            if m:
                f["size"] = float(m.group(1)) * (0.75 if m.group(2) == "px" else 1)
            if n == "br":
                para.add_run().add_break()
            elif n == "img":
                img = _data_image(ch.get("src", ""))
                if img:
                    try:
                        para.add_run().add_picture(img, width=Inches(min(6.0, float(ch.get("width") or 0) / 96 or 6.0)))
                    except Exception:
                        pass
            else:
                runs(ch, para, f)

    def block(node, list_style=None):
        for ch in node.children:
            if isinstance(ch, NavigableString):
                if ch.strip():
                    p = doc.add_paragraph()
                    p.add_run(re.sub(r"\s+", " ", str(ch)))
                continue
            if not isinstance(ch, Tag):
                continue
            n = ch.name
            if n in ("h1", "h2", "h3", "h4", "h5", "h6"):
                p = doc.add_heading(level=int(n[1]))
                runs(ch, p, {})
            elif n in ("p", "pre", "blockquote", "address"):
                p = doc.add_paragraph(style="Quote" if n == "blockquote" else None)
                runs(ch, p, {"mono": n == "pre"})
                al = _style(ch).get("text-align") or ch.get("align", "")
                if al in ALIGN:
                    p.alignment = ALIGN[al]
            elif n in ("ul", "ol"):
                for li in ch.find_all("li", recursive=False):
                    p = doc.add_paragraph(style="List Number" if n == "ol" else "List Bullet")
                    inner = [c for c in li.children if not (isinstance(c, Tag) and c.name in ("ul", "ol"))]
                    holder = BeautifulSoup("", "html.parser").new_tag("span")
                    for c in inner:
                        holder.append(c.__copy__() if isinstance(c, Tag) else NavigableString(str(c)))
                    runs(holder, p, {})
                    for sub in li.find_all(["ul", "ol"], recursive=False):
                        block(BeautifulSoup(str(sub), "html.parser"))
            elif n == "table":
                rows = ch.find_all("tr")
                if not rows:
                    continue
                cols = max(len(r.find_all(["td", "th"])) for r in rows) or 1
                t = doc.add_table(rows=len(rows), cols=cols)
                t.style = "Table Grid"
                for i, r in enumerate(rows):
                    for j, c in enumerate(r.find_all(["td", "th"])[:cols]):
                        cell = t.cell(i, j)
                        cell.paragraphs[0].text = ""
                        runs(c, cell.paragraphs[0], {"b": c.name == "th"})
            elif n == "hr":
                doc.add_paragraph("―" * 20).alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif n == "img":
                p = doc.add_paragraph()
                runs(BeautifulSoup(str(ch), "html.parser"), p, {})
            elif n in ("div", "section", "article", "main", "body", "html", "header", "footer", "center", "figure"):
                block(ch)
            else:
                p = doc.add_paragraph()
                runs(ch, p, {})

    block(soup)
    doc.save(str(out))
    return out


def _font_files() -> dict:
    """Шрифт с кириллицей для PDF: Android, Linux, Windows, macOS."""
    cands = [
        ("/system/fonts/Roboto-Regular.ttf", "/system/fonts/Roboto-Bold.ttf", "/system/fonts/Roboto-Italic.ttf", "/system/fonts/Roboto-BoldItalic.ttf"),
        ("/system/fonts/NotoSans-Regular.ttf", "/system/fonts/NotoSans-Bold.ttf", "/system/fonts/NotoSans-Italic.ttf", "/system/fonts/NotoSans-BoldItalic.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf"),
        (r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\arialbd.ttf", r"C:\Windows\Fonts\ariali.ttf", r"C:\Windows\Fonts\arialbi.ttf"),
        ("/Library/Fonts/Arial Unicode.ttf",) * 4,
    ]
    for reg, bold, it, bi in cands:
        if os.path.exists(reg):
            pick = lambda p: p if os.path.exists(p) else reg   # noqa: E731
            return {"": reg, "B": pick(bold), "I": pick(it), "BI": pick(bi)}
    raise DocError("нет шрифта с кириллицей для PDF")


def _clean_for_pdf(html: str) -> str:
    """fpdf2 понимает простой HTML: убираем стили и незнакомые теги, оставляя текст."""
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    keep = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "b", "strong", "i", "em", "u", "a", "ul", "ol", "li", "br",
            "table", "thead", "tbody", "tr", "td", "th", "img", "hr", "blockquote", "pre", "code", "sup", "sub"}
    for t in soup.find_all(True):
        if t.name not in keep:
            t.unwrap()
            continue
        attrs = {k: v for k, v in t.attrs.items() if k in ("href", "src", "width", "height", "colspan")}
        if t.name == "img" and not str(attrs.get("src", "")).startswith("data:"):
            t.decompose()
            continue
        t.attrs = attrs
    for t in soup.find_all(["td", "th"]):       # таблицы fpdf2: внутри ячеек только текст
        t.string = t.get_text(" ", strip=True)
    return str(soup)


def to_pdf(html: str, out: Path, title: str = "") -> Path:
    from fpdf import FPDF
    pdf = FPDF(format="A4")
    pdf.set_margins(20, 18, 20)
    pdf.set_auto_page_break(True, 18)
    if title:
        pdf.set_title(title)
    for style, path in _font_files().items():
        pdf.add_font("Uni", style, path)
    pdf.add_page()
    pdf.set_font("Uni", size=12)
    try:
        try:   # код и pre тоже шрифтом с кириллицей (по умолчанию Courier без неё)
            from fpdf.fonts import FontFace
            styles = {"code": FontFace(family="Uni"), "pre": FontFace(family="Uni")}
            pdf.write_html(_clean_for_pdf(html), font_family="Uni", tag_styles=styles)
        except TypeError:          # старый fpdf2 без tag_styles
            pdf.write_html(_clean_for_pdf(html), font_family="Uni", pre_code_font="Uni")
    except Exception:
        # сложный HTML: хотя бы текст абзацами
        pdf = FPDF(format="A4")
        pdf.set_margins(20, 18, 20)
        for style, path in _font_files().items():
            pdf.add_font("Uni", style, path)
        pdf.add_page()
        pdf.set_font("Uni", size=12)
        pdf.multi_cell(0, 6, to_text(html))
    pdf.output(str(out))
    return out


def write(html: str, fmt: str, out: Path, title: str = "") -> Path:
    if fmt == "html":
        page = (f'<!doctype html><html lang="ru"><head><meta charset="utf-8"><title>{htmllib.escape(title)}</title>'
                f"</head><body>{html}</body></html>")
        out.write_text(page, encoding="utf-8")
    elif fmt == "md":
        out.write_text(to_markdown(html), encoding="utf-8")
    elif fmt == "txt":
        out.write_text(to_text(html), encoding="utf-8")
    elif fmt == "docx":
        to_docx(html, out, title)
    elif fmt == "pdf":
        to_pdf(html, out, title)
    else:
        raise DocError(f"встроенный конвертер не пишет .{fmt}")
    return out


def convert(src: Path, out: Path, to: str) -> Path:
    """Файл -> файл (для «Документов»)."""
    ext = src.suffix.lower().lstrip(".")
    if ext not in READ:
        raise DocError(f"на телефоне .{ext} не открывается (нужен pandoc)")
    if to not in WRITE and to != "pdf":
        raise DocError(f"на телефоне нельзя сделать .{to}")
    return write(read_html(src), to, out, src.stem)
