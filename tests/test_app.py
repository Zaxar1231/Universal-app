"""
Проверки на реальных файлах: ffmpeg генерирует тестовые ролики, дальше
смотрим ffprobe-ом, что получилось именно то, что просили.

Запуск: python -m pytest -q
Тесты, которым нужен интернет (yt-dlp), включаются переменной NET=1.
"""
import io
import json
import os
import shutil
import subprocess
import sys
import time
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import app as webapp  # noqa: E402
import downloader  # noqa: E402
import editor  # noqa: E402
import qrtools  # noqa: E402
from converters import ConvertError  # noqa: E402

needs_ffmpeg = pytest.mark.skipif(not shutil.which("ffmpeg"), reason="нет ffmpeg")


def ff(*args):
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", *args], check=True)


def streams(path):
    res = subprocess.run(["ffprobe", "-v", "error", "-print_format", "json", "-show_streams",
                          "-show_format", str(path)], capture_output=True, text=True, check=True)
    return json.loads(res.stdout)


@pytest.fixture(scope="session")
def media(tmp_path_factory):
    d = tmp_path_factory.mktemp("media")
    if not shutil.which("ffmpeg"):
        return d
    ff("-f", "lavfi", "-i", "testsrc2=s=640x360:r=30:d=6", "-f", "lavfi", "-i", "sine=f=440:d=6",
       "-c:v", "libx264", "-g", "30", "-sc_threshold", "0", "-c:a", "aac", "-shortest", str(d / "a.mp4"))
    ff("-f", "lavfi", "-i", "testsrc=s=320x240:r=25:d=4", "-c:v", "libx264", str(d / "silent.mp4"))
    ff("-f", "lavfi", "-i", "sine=f=330:d=5", str(d / "tone.mp3"))
    return d


@pytest.fixture()
def client():
    webapp.app.config["TESTING"] = True
    return webapp.app.test_client()


def wait_task(client, tid, timeout=120):
    t0 = time.time()
    while time.time() - t0 < timeout:
        d = client.get(f"/api/task/{tid}").get_json()
        if d["status"] in ("done", "error", "cancelled"):
            return d
        time.sleep(0.2)
    raise AssertionError("задача не завершилась")


# ---------------------------------------------------------------- редактор

def test_parse_time():
    assert editor.parse_time("90") == 90
    assert editor.parse_time("1:30") == 90
    assert editor.parse_time("01:02:03.5") == 3723.5
    assert editor.parse_time("1м30с") == 90
    assert editor.parse_time("2,5") == 2.5
    assert editor.parse_time("") is None
    with pytest.raises(ConvertError):
        editor.parse_time("abc")


def test_atempo_chain():
    assert editor._atempo_chain(4) == ["atempo=2.0", "atempo=2"]
    assert editor._atempo_chain(0.25) == ["atempo=0.5", "atempo=0.5"]


def opts(**kw):
    base = {k: "" for k in ("start", "end")}
    base.update({k: str(v) for k, v in kw.items()})
    return editor.read_opts(base)


@needs_ffmpeg
def test_trim_crop_scale(media, tmp_path):
    out = editor.edit_file(media / "a.mp4", tmp_path, "r",
                           opts(start="1", end="4", aspect="9:16", height="240", out="mp4"))
    s = streams(out)
    v = next(x for x in s["streams"] if x["codec_type"] == "video")
    assert v["height"] == 240 and v["width"] in (134, 136)          # 9:16 после кропа
    assert abs(float(s["format"]["duration"]) - 3) < 0.25
    assert any(x["codec_type"] == "audio" for x in s["streams"])


@needs_ffmpeg
def test_speed_and_mute(media, tmp_path):
    out = editor.edit_file(media / "a.mp4", tmp_path, "r", opts(speed="2", mute="1", out="webm"))
    s = streams(out)
    assert abs(float(s["format"]["duration"]) - 3) < 0.3
    assert [x["codec_name"] for x in s["streams"]] == ["vp9"]


@needs_ffmpeg
def test_target_size(media, tmp_path):
    out = editor.edit_file(media / "a.mp4", tmp_path, "r", opts(target_mb="0.3", out="mp4"))
    assert out.stat().st_size < 0.3 * 1048576 * 1.15


@needs_ffmpeg
def test_target_size_too_small(media, tmp_path):
    with pytest.raises(ConvertError, match="слишком мало"):
        editor.edit_file(media / "a.mp4", tmp_path, "r", opts(target_mb="0.001", out="mp4"))


@needs_ffmpeg
def test_gif_and_rotate(media, tmp_path):
    out = editor.edit_file(media / "a.mp4", tmp_path, "r", opts(out="gif", rotate="90", gif_width="200", end="2"))
    v = streams(out)["streams"][0]
    assert v["codec_name"] == "gif" and v["width"] == 200 and v["height"] > 200


@needs_ffmpeg
def test_audio_edit(media, tmp_path):
    out = editor.edit_file(media / "tone.mp3", tmp_path, "r",
                           opts(out="same", fade_in="1", fade_out="1", volume="50", normalize="1", end="3"))
    s = streams(out)
    assert out.suffix == ".mp3" and abs(float(s["format"]["duration"]) - 3) < 0.2


@needs_ffmpeg
def test_audio_file_to_video_format_rejected(media, tmp_path):
    with pytest.raises(ConvertError, match="нет видео"):
        editor.edit_file(media / "tone.mp3", tmp_path, "r", opts(out="mp4"))


@needs_ffmpeg
def test_extract_audio(media, tmp_path):
    out = editor.edit_file(media / "a.mp4", tmp_path, "r", opts(out="opus"))
    assert [x["codec_name"] for x in streams(out)["streams"]] == ["opus"]


@needs_ffmpeg
def test_fast_trim_copy(media, tmp_path):
    out = editor.edit_file(media / "a.mp4", tmp_path, "r", opts(start="2", fast="1", out="same"))
    # без перекодирования режется по ключевому кадру (каждую секунду): свежий ffmpeg берёт кадр строго
    # до точки (1 с, выходит 5 с), ffmpeg 6.x — на ней (2 с, выходит 4 с). Исходник 6 с.
    assert out.suffix == ".mp4" and 3.5 < float(streams(out)["format"]["duration"]) < 5.5


@needs_ffmpeg
def test_merge_mixed_sizes_and_silence(media, tmp_path):
    out = editor.merge_files([media / "a.mp4", media / "silent.mp4"], tmp_path, "m", opts(out="mp4"))
    s = streams(out)
    v = next(x for x in s["streams"] if x["codec_type"] == "video")
    assert (v["width"], v["height"]) == (640, 360)
    assert any(x["codec_type"] == "audio" for x in s["streams"])
    assert abs(float(s["format"]["duration"]) - 10) < 0.4


@needs_ffmpeg
def test_edit_endpoint_and_cancel(client, media):
    with open(media / "a.mp4", "rb") as f:
        r = client.post("/api/edit", data={"files": (f, "клип.mp4"), "out": "mp4", "height": "240"},
                        content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    d = wait_task(client, r.get_json()["task"])
    assert d["status"] == "done" and d["single_file"] == "клип_edit.mp4"
    dl = client.get(f"/api/download/{d['job']}?raw=1")
    assert dl.status_code == 200 and dl.data[4:8] == b"ftyp"
    assert "UTF-8''%D0%BA%D0%BB%D0%B8%D0%BF_edit.mp4" in dl.headers["Content-Disposition"]
    dl.close()

    # отмена: длинная задача (медленный пресет) убивается
    with open(media / "a.mp4", "rb") as f:
        r = client.post("/api/edit", data={"files": (f, "a.mp4"), "out": "mkv", "vcodec": "h265",
                                           "preset": "veryslow", "speed": "0.25"},
                        content_type="multipart/form-data")
    tid = r.get_json()["task"]
    time.sleep(0.8)
    client.post(f"/api/task/{tid}/cancel")
    assert wait_task(client, tid)["status"] == "cancelled"


def test_edit_endpoint_rejects_non_media(client):
    r = client.post("/api/edit", data={"files": (io.BytesIO(b"x"), "doc.txt")},
                    content_type="multipart/form-data")
    assert r.status_code == 400


# ---------------------------------------------------------------- QR

def test_qr_png_decodes_size(client):
    r = client.post("/api/qr", data={"text": "https://example.com", "box_size": "5", "border": "4"})
    assert r.status_code == 200 and r.mimetype == "image/png"
    from PIL import Image
    im = Image.open(io.BytesIO(r.data))
    n = len(qrtools.matrix("https://example.com", "M"))
    assert im.size == ((n + 8) * 5,) * 2
    assert im.getpixel((0, 0)) == (255, 255, 255)
    assert im.getpixel((20, 20)) == (0, 0, 0)            # угол поискового узора


def test_qr_svg_colors_transparent(client):
    r = client.post("/api/qr", data={"text": "hi", "format": "svg", "fg": "#ff0000", "transparent": "1"})
    body = r.data.decode()
    assert r.mimetype == "image/svg+xml" and 'fill="#ff0000"' in body and "<rect" not in body


def test_qr_wifi_payload():
    assert qrtools.wifi_payload("My;Net", "p:ss") == r"WIFI:T:WPA;S:My\;Net;P:p\:ss;;"
    assert qrtools.wifi_payload("Open", "", "nopass") == "WIFI:T:nopass;S:Open;;"


def test_qr_bad_input(client):
    assert client.post("/api/qr", data={"text": " "}).status_code == 400
    assert client.post("/api/qr", data={"text": "x", "fg": "red"}).status_code == 400
    assert client.post("/api/qr", data={"text": "x" * 5000}).status_code == 400


def make_docx(links):
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rels = "".join(f'<Relationship Id="r{i}" Type="hyperlink" Target="{u}" TargetMode="External"/>'
                   for i, u in enumerate(links))
    body = "".join(f'<w:p><w:hyperlink r:id="r{i}"><w:r><w:t>link</w:t></w:r></w:hyperlink></w:p>'
                   for i in range(len(links)))
    body += '<w:p><w:r><w:t>see https://plain.example/x</w:t></w:r></w:p>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("word/_rels/document.xml.rels",
                   f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>')
        z.writestr("word/document.xml", f'<w:document xmlns:w="{W}" xmlns:r="{R}"><w:body>{body}</w:body></w:document>')
    return buf.getvalue()


def test_qr_bulk_sources(client):
    data = {
        "links": "https://a.example\n# комментарий\nhttps://a.example\n",
        "column": "url",
        "files": [(io.BytesIO(b"name;url\nx;https://b.example\ny;https://c.example\n"), "list.csv"),
                  (io.BytesIO(make_docx(["https://d.example/1"])), "doc.docx")],
        "format": "svg",
    }
    r = client.post("/api/qr/bulk", data=data, content_type="multipart/form-data")
    d = r.get_json()
    assert r.status_code == 200, d
    assert d["total"] == 5 and d["done"] == 5                  # дубль убран
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/download/{d['job']}").data))
    names = sorted(z.namelist())
    assert "index.csv" in names and "0001_a.example.svg" in names and "0005_plain.example.svg" in names


def test_qr_bulk_wrong_column(client):
    r = client.post("/api/qr/bulk", data={"column": "link", "files": [(io.BytesIO(b"url\nhttps://x"), "a.csv")]},
                    content_type="multipart/form-data")
    assert r.status_code == 400 and "колонки" in r.get_json()["error"]


# ---------------------------------------------------------------- загрузчик

def test_extract_urls():
    text = "смотри https://youtu.be/abc, и вот (https://vk.com/video1_2). https://youtu.be/abc"
    assert downloader.extract_urls(text) == ["https://youtu.be/abc", "https://vk.com/video1_2"]


def test_dl_opts_audio():
    o = downloader.read_opts({"mode": "audio", "audio_format": "mp3", "abitrate": "192"})
    y = downloader.build_opts(o, Path("/tmp"), None)
    pp = y["postprocessors"][0]
    assert y["format"] == "bestaudio/best" and pp["preferredcodec"] == "mp3" and pp["preferredquality"] == "192"
    assert y["noplaylist"] is True


def test_dl_opts_video_quality():
    o = downloader.read_opts({"mode": "video", "quality": "720", "container": "mp4", "playlist": "1"})
    y = downloader.build_opts(o, Path("/tmp"), None)
    assert "[height<=720]" in y["format"] and y["format_sort"][0] == "res:720"
    assert y["merge_output_format"] == "mp4" and y["noplaylist"] is False


def test_dl_rejects_no_urls(client):
    assert client.post("/api/dl", data={"urls": "просто текст"}).status_code == 400


@pytest.fixture()
def http_media(media):
    """Локальный HTTP-сервер: yt-dlp качает прямую ссылку generic-экстрактором."""
    import functools
    import http.server
    import threading
    handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(media))
    handler.log_message = lambda *a: None
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_port}"
    srv.shutdown()


@needs_ffmpeg
def test_dl_local_video_and_audio(client, http_media):
    r = client.post("/api/dl", data={"urls": f"{http_media}/a.mp4", "mode": "video"})
    d = wait_task(client, r.get_json()["task"])
    assert d["status"] == "done" and d["single_file"].endswith(".mp4"), d

    r = client.post("/api/dl", data={"urls": f"{http_media}/a.mp4 {http_media}/nope.mp4",
                                     "mode": "audio", "audio_format": "mp3", "abitrate": "128"})
    d = wait_task(client, r.get_json()["task"])
    assert d["status"] == "done" and d["done"] == 1 and len(d["failed"]) == 1, d
    assert d["single_file"].endswith(".mp3")


@needs_ffmpeg
def test_dl_info_local(client, http_media):
    d = client.post("/api/dl/info", data={"urls": f"{http_media}/a.mp4"}).get_json()
    assert d["items"][0]["title"] == "a" and not d["items"][0].get("error")


@pytest.mark.skipif(os.environ.get("NET") != "1", reason="нужен интернет: NET=1")
def test_dl_real(client):
    url = "https://www.youtube.com/watch?v=jNQXAC9IVRw"   # «Me at the zoo», 19 секунд
    r = client.post("/api/dl", data={"urls": url, "mode": "audio", "audio_format": "mp3"})
    d = wait_task(client, r.get_json()["task"], timeout=300)
    assert d["status"] == "done", d


# ---------------------------------------------------------------- конвертер (регресс)

@needs_ffmpeg
def test_universal_still_works(client, media):
    with open(media / "tone.mp3", "rb") as f:
        r = client.post("/api/universal", data={"files": (f, "tone.mp3"), "to": "flac"},
                        content_type="multipart/form-data")
    d = wait_task(client, r.get_json()["task"])
    assert d["status"] == "done" and d["done"] == 1 and d["single_file"] == "tone.flac"


@pytest.mark.skipif(not shutil.which("convert"), reason="нет ImageMagick")
def test_universal_image_resize(client):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (1200, 600), (10, 200, 30)).save(buf, "PNG")
    buf.seek(0)
    r = client.post("/api/universal", data={"files": (buf, "big.png"), "to": "jpg", "resize": "300"},
                    content_type="multipart/form-data")
    d = wait_task(client, r.get_json()["task"])
    im = Image.open(io.BytesIO(client.get(f"/api/download/{d['job']}?raw=1").data))
    assert im.size == (300, 150)


# ---------------------------------------------------------------- PDF

def make_pdf(pages=3, photo=False):
    from PIL import Image
    import random
    imgs = []
    for i in range(pages):
        im = Image.new("RGB", (1600, 2200) if photo else (200, 300), (255, 255, 255))
        if photo:   # шум плохо жмётся без потерь: имитация скана
            im = Image.frombytes("RGB", (1600, 2200), bytes(random.getrandbits(8) for _ in range(1600 * 2200 * 3)))
        imgs.append(im)
    buf = io.BytesIO()
    imgs[0].save(buf, "PDF", save_all=True, append_images=imgs[1:])
    return buf.getvalue()


def pdf_task(client, files, **form):
    data = {"files": [(io.BytesIO(b), n) for n, b in files], **form}
    r = client.post("/api/pdf", data=data, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    return wait_task(client, r.get_json()["task"])


def pages_of(client, d, name=None):
    from pypdf import PdfReader
    url = f"/api/file/{d['job']}/{name}" if name else f"/api/download/{d['job']}?raw=1"
    return len(PdfReader(io.BytesIO(client.get(url).data)).pages)


def test_pdf_ranges():
    import pdftools
    assert pdftools.parse_ranges("1-3, 5, 8-", 9) == [0, 1, 2, 4, 7, 8]
    assert pdftools.parse_ranges("", 2) == [0, 1]
    with pytest.raises(pdftools.PdfError):
        pdftools.parse_ranges("4", 3)
    with pytest.raises(pdftools.PdfError):
        pdftools.parse_ranges("a-b", 3)


def test_pdf_merge_split_extract_rotate(client):
    a, b = make_pdf(3), make_pdf(2)
    d = pdf_task(client, [("a.pdf", a), ("b.pdf", b)], op="merge", merge_name="all")
    assert d["single_file"] == "all.pdf" and pages_of(client, d) == 5

    d = pdf_task(client, [("a.pdf", a)], op="split")
    assert d["names"] == ["a_p1.pdf", "a_p2.pdf", "a_p3.pdf"]

    d = pdf_task(client, [("a.pdf", a)], op="extract", ranges="3,1")
    assert pages_of(client, d) == 2

    from pypdf import PdfReader
    d = pdf_task(client, [("a.pdf", a)], op="rotate", angle="90", ranges="2")
    r = PdfReader(io.BytesIO(client.get(f"/api/download/{d['job']}?raw=1").data))
    assert [p.rotation for p in r.pages] == [0, 90, 0]


def test_pdf_compress(client):
    src = make_pdf(1, photo=True)
    d = pdf_task(client, [("scan.pdf", src)], op="compress", quality="50", max_px="800")
    out = client.get(f"/api/download/{d['job']}?raw=1").data
    assert d["status"] == "done" and len(out) < len(src) / 3


def test_pdf_password_roundtrip(client):
    d = pdf_task(client, [("a.pdf", make_pdf(2))], op="protect", new_password="тайна1")
    locked = client.get(f"/api/download/{d['job']}?raw=1").data
    from pypdf import PdfReader
    assert PdfReader(io.BytesIO(locked)).is_encrypted

    d = pdf_task(client, [("x.pdf", locked)], op="rotate", angle="180")
    assert d["status"] == "error" and "паролем" in d["failed"][0]["error"]
    d = pdf_task(client, [("x.pdf", locked)], op="unlock", password="не тот")
    assert "Неверный пароль" in d["failed"][0]["error"]
    d = pdf_task(client, [("x.pdf", locked)], op="unlock", password="тайна1")
    assert not PdfReader(io.BytesIO(client.get(f"/api/download/{d['job']}?raw=1").data)).is_encrypted


def test_pdf_bad_requests(client):
    one = [("a.pdf", make_pdf(1))]
    r = client.post("/api/pdf", data={"op": "merge", "files": [(io.BytesIO(b), n) for n, b in one]},
                    content_type="multipart/form-data")
    assert r.status_code == 400
    d = pdf_task(client, [("bad.pdf", b"not a pdf")], op="rotate")
    assert d["status"] == "error"


def test_index_renders(client):
    html = client.get("/").data.decode()
    for mode in ("image", "edit", "dl", "qr", "topdf", "pdft"):
        assert f'value="{mode}"' in html


def test_password(monkeypatch, client):
    monkeypatch.setattr(webapp, "APP_PASSWORD", "secret")
    assert client.get("/api/engines").status_code == 401
    import base64
    ok = {"Authorization": "Basic " + base64.b64encode(b"me:secret").decode()}
    assert client.get("/api/engines", headers=ok).status_code == 200


# ---------------------------------------------------------------- движки (настольная версия)

@pytest.fixture()
def engine_env(tmp_path, monkeypatch):
    import engines
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path / "home"))
    monkeypatch.setenv("UNIVERSAL_ENGINES_ANYOS", "1")
    srv_dir = tmp_path / "srv"
    srv_dir.mkdir()
    import functools
    import http.server
    import threading
    handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=str(srv_dir))
    handler.log_message = lambda *a: None
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield engines, srv_dir, f"http://127.0.0.1:{srv.server_port}"
    srv.shutdown()


def test_engine_install_and_activate(engine_env, monkeypatch, client):
    engines, srv_dir, base = engine_env
    with zipfile.ZipFile(srv_dir / "ff.zip", "w") as z:
        z.writestr("ffmpeg-build/bin/ffmpeg.exe", "x")
        z.writestr("ffmpeg-build/bin/ffprobe.exe", "x")
        z.writestr("ffmpeg-build/bin/avcodec.dll", "x")
    monkeypatch.setitem(engines.CATALOG, "ffmpeg", {**engines.CATALOG["ffmpeg"], "url": base + "/ff.zip"})
    monkeypatch.setenv("PATH", "/usr/bin:/bin")

    r = client.post("/api/engines/install/ffmpeg")
    d = wait_task(client, r.get_json()["task"])
    assert d["status"] == "done", d
    bindir = engines.root() / "engines" / "ffmpeg" / "ffmpeg-build" / "bin"
    assert (bindir / "avcodec.dll").exists()                      # dll рядом с exe
    assert os.environ["PATH"].split(os.pathsep)[0] == str(bindir)
    assert engines.status()["ffmpeg"]["installed"]

    monkeypatch.setenv("PATH", "/usr/bin:/bin")
    engines.activate()                                            # повторный запуск приложения
    assert str(bindir) in os.environ["PATH"]


def test_engine_install_errors(engine_env, monkeypatch, client):
    engines, srv_dir, base = engine_env
    with zipfile.ZipFile(srv_dir / "bad.zip", "w") as z:
        z.writestr("readme.txt", "no exe here")
    with zipfile.ZipFile(srv_dir / "evil.zip", "w") as z:
        z.writestr("../../evil/pandoc.exe", "x")
    monkeypatch.setitem(engines.CATALOG, "pandoc", {**engines.CATALOG["pandoc"], "url": base + "/bad.zip"})
    d = wait_task(client, client.post("/api/engines/install/pandoc").get_json()["task"])
    assert d["status"] == "error" and "нет pandoc.exe" in d["error"]

    monkeypatch.setitem(engines.CATALOG, "pandoc", {**engines.CATALOG["pandoc"], "url": base + "/evil.zip"})
    d = wait_task(client, client.post("/api/engines/install/pandoc").get_json()["task"])
    assert d["status"] == "error" and "повреждён" in d["error"]
    assert not (engines.root().parent / "evil").exists()

    monkeypatch.setitem(engines.CATALOG, "pandoc", {**engines.CATALOG["pandoc"], "url": base + "/missing.zip"})
    d = wait_task(client, client.post("/api/engines/install/pandoc").get_json()["task"])
    assert d["status"] == "error" and "прервалась" in d["error"]
    assert client.post("/api/engines/install/nope").status_code == 404


def test_engine_ytdlp_zipapp(engine_env, monkeypatch):
    engines, srv_dir, base = engine_env
    with zipfile.ZipFile(srv_dir / "yt-dlp", "w") as z:
        z.writestr("fake_ytdlp_probe/__init__.py", "VERSION = 'fresh'")
    monkeypatch.setitem(engines.CATALOG, "ytdlp", {**engines.CATALOG["ytdlp"], "url": base + "/yt-dlp"})
    path = engines.install("ytdlp")
    monkeypatch.setattr(sys, "path", list(sys.path))
    engines.activate()
    assert sys.path[0] == path
    import importlib
    assert importlib.import_module("fake_ytdlp_probe").VERSION == "fresh"


def test_engine_install_only_on_windows(monkeypatch, client):
    monkeypatch.delenv("UNIVERSAL_ENGINES_ANYOS", raising=False)
    if sys.platform != "win32":
        assert client.post("/api/engines/install/ffmpeg").status_code == 400


def test_windows_ignores_system_convert(monkeypatch):
    import converters
    monkeypatch.setattr(converters.sys, "platform", "win32")
    monkeypatch.setattr(converters.shutil, "which", lambda n: "C:/Windows/System32/convert.exe" if n == "convert" else None)
    assert converters._magick() is None


def test_engine_resolve_falls_back_to_release_page(monkeypatch):
    """API упёрся в лимит (403) -> ссылка берётся со страницы релиза."""
    import http.server
    import threading
    import engines

    class FakeGitHub(http.server.BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_GET(self):
            if self.path.startswith("/api/"):
                self.send_response(403); self.end_headers(); self.wfile.write(b"rate limit exceeded")
            elif self.path == "/jgm/pandoc/releases/latest":
                self.send_response(302); self.send_header("Location", "/jgm/pandoc/releases/tag/3.9"); self.end_headers()
            elif self.path == "/jgm/pandoc/releases/tag/3.9":
                self.send_response(200); self.end_headers()
            elif self.path == "/jgm/pandoc/releases/expanded_assets/3.9":
                self.send_response(200); self.end_headers()
                self.wfile.write(b'<a href="/jgm/pandoc/releases/download/3.9/pandoc-3.9-arm64-macOS.zip">'
                                 b'<a href="/jgm/pandoc/releases/download/3.9/pandoc-3.9-windows-x86_64.zip">')
            else:
                self.send_response(404); self.end_headers()

    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), FakeGitHub)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_port}"
    monkeypatch.setattr(engines, "GITHUB_API", base + "/api/repos/{repo}/releases/latest")
    monkeypatch.setattr(engines, "GITHUB_WEB", base)
    try:
        assert engines._resolve_url(engines.CATALOG["pandoc"]) == base + "/jgm/pandoc/releases/download/3.9/pandoc-3.9-windows-x86_64.zip"
        with pytest.raises(engines.EngineError, match="нет файла"):
            engines._resolve_url({"repo": "jgm/pandoc", "asset": r"nothing-matches$"})
        with pytest.raises(engines.EngineError, match="Проверь интернет"):
            engines._resolve_url({"repo": "no/such", "asset": "x"})
    finally:
        srv.shutdown()


def test_engine_7z_falls_back_to_7zip(engine_env, monkeypatch, tmp_path):
    """py7zr не осилил архив (BCJ2) -> распаковка внешним 7-Zip."""
    engines, srv_dir, base = engine_env
    py7zr = pytest.importorskip("py7zr")      # есть в настольных зависимостях, в обычных нет
    with py7zr.SevenZipFile(srv_dir / "im.7z", "w") as z:
        z.writestr("x", "ImageMagick-7/magick.exe")
    calls = []

    class Broken:
        def __init__(self, *a, **k):
            raise RuntimeError("BCJ2 filter is not supported by py7zr")

    def fake_run(cmd, **kw):
        calls.append(cmd)
        out = Path([c for c in cmd if c.startswith("-o")][0][2:])
        (out / "ImageMagick-7").mkdir(parents=True)
        (out / "ImageMagick-7" / "magick.exe").write_text("x")
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(py7zr, "SevenZipFile", Broken)
    monkeypatch.setattr(engines, "_seven_zip_exe", lambda: "7zr")
    monkeypatch.setattr(subprocess, "run", fake_run)
    monkeypatch.setitem(engines.CATALOG, "imagemagick", {**engines.CATALOG["imagemagick"], "url": base + "/im.7z"})
    path = engines.install("imagemagick")
    assert calls and calls[0][:2] == ["7zr", "x"]
    assert Path(path, "magick.exe").exists()


def test_engine_asset_priority():
    import engines
    spec = engines.CATALOG["imagemagick"]
    assets = [("ImageMagick-7.1.2-portable-Q16-HDRI-x64.7z", "u7z"), ("ImageMagick-7.1.2-portable-Q16-HDRI-x64.zip", "uzip")]
    orig = engines._assets_from_api
    try:
        engines._assets_from_api = lambda repo: assets
        assert engines._resolve_url(spec) == "uzip"
        engines._assets_from_api = lambda repo: assets[:1]
        assert engines._resolve_url(spec) == "u7z"
    finally:
        engines._assets_from_api = orig


# ---------------------------------------------------------------- QR в документе

def make_qr_docx() -> bytes:
    """Word-файл со всеми видами ссылок: гиперссылка, поле HYPERLINK, адрес текстом."""
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    body = (
        '<w:p><w:r><w:t xml:space="preserve">Сайт: </w:t></w:r>'
        '<w:hyperlink r:id="rId5"><w:r><w:t>наш сайт</w:t></w:r></w:hyperlink>'
        '<w:r><w:t>.</w:t></w:r></w:p>'
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText xml:space="preserve"> HYPERLINK "https://field.example.com/x" </w:instrText></w:r>'
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>поле</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
        '<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Текстом https://plain.example.com/a?b=1&amp;c=2, и дальше.</w:t></w:r></w:p>'
        '<w:p><w:hyperlink r:id="rId6"><w:r><w:t>почта</w:t></w:r></w:hyperlink></w:p>'
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml",
                   '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr("word/_rels/document.xml.rels",
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://site.example.com/" TargetMode="External"/>'
                   '<Relationship Id="rId6" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="mailto:a@b.c" TargetMode="External"/>'
                   '</Relationships>')
        z.writestr("word/document.xml", f'<w:document xmlns:w="{W}" xmlns:r="{R}"><w:body>{body}</w:body></w:document>')
    return buf.getvalue()


def test_qrdoc_docx_scan_and_replace(tmp_path):
    import qrdoc
    from xml.etree import ElementTree as ET
    src = tmp_path / "doc.docx"
    src.write_bytes(make_qr_docx())
    links = {l["url"]: l for l in qrdoc.scan(src, tmp_path / "w")["links"]}
    assert set(links) == {"https://site.example.com/", "https://field.example.com/x", "https://plain.example.com/a?b=1&c=2"}
    assert links["https://site.example.com/"]["text"] == "наш сайт" and links["https://field.example.com/x"]["text"] == "поле"

    out = tmp_path / "out"; out.mkdir()
    st = qrtools.read_style({})
    dst, n = qrdoc.apply(src, out, "doc", set(links) - {"https://field.example.com/x"}, "replace", 3, st, tmp_path / "w")
    assert n == 2
    z = zipfile.ZipFile(dst)
    xml = z.read("word/document.xml").decode()
    ET.fromstring(xml)                                           # XML не сломан
    assert xml.count("<w:drawing>") == 2
    assert "наш сайт" not in xml and "HYPERLINK" in xml           # гиперссылка заменена, поле оставлено
    import re as _re
    texts = "".join(_re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))
    assert "Текстом , и дальше." in texts and "plain.example.com" not in texts   # адрес ушёл, текст вокруг остался
    assert '<w:b/>' in xml                                       # форматирование куска текста сохранилось
    assert "mailto" in z.read("word/_rels/document.xml.rels").decode()
    assert 'Extension="png"' in z.read("[Content_Types].xml").decode()
    pngs = [n for n in z.namelist() if n.startswith("word/media/")]
    assert len(pngs) == 2 and all(z.read(p)[:4] == b"\x89PNG" for p in pngs)
    assert [l["url"] for l in qrdoc.scan(dst, tmp_path / "w2")["links"]] == ["https://field.example.com/x"]


def test_qrdoc_docx_after_keeps_links(tmp_path):
    import qrdoc
    src = tmp_path / "doc.docx"
    src.write_bytes(make_qr_docx())
    out = tmp_path / "out"; out.mkdir()
    dst, n = qrdoc.apply(src, out, "doc", {"https://field.example.com/x", "https://plain.example.com/a?b=1&c=2"},
                         "after", 2, qrtools.read_style({}), tmp_path / "w")
    xml = zipfile.ZipFile(dst).read("word/document.xml").decode()
    assert n == 2 and xml.count("<w:drawing>") == 2
    assert "https://plain.example.com/a?b=1&amp;c=2</w:t>" in xml and "поле" in xml
    assert len(qrdoc.scan(dst, tmp_path / "w2")["links"]) == 3   # все ссылки на месте


def test_qrdoc_markdown(tmp_path):
    import qrdoc
    md = ("# Заголовок\n\n[Сайт](https://a.example.com/x \"title\") и <https://b.example.com>.\n\n"
          "Голая https://c.example.com/p, дальше. ![картинка](https://img.example.com/i.png)\n\n"
          "```\nhttps://code.example.com\n```\n\nКод `https://inline.example.com` тоже.\n")
    src = tmp_path / "note.md"
    src.write_text(md, encoding="utf-8")
    urls = [l["url"] for l in qrdoc.scan(src, tmp_path)["links"]]
    assert urls == ["https://a.example.com/x", "https://b.example.com", "https://c.example.com/p"]
    out = tmp_path / "out"; out.mkdir()
    dst, n = qrdoc.apply(src, out, "note", set(urls), "replace", 3, qrtools.read_style({}), tmp_path / "w")
    text = dst.read_text(encoding="utf-8")
    assert n == 3 and text.count('<img src="note_qr_') == 3
    assert "https://code.example.com" in text and "https://inline.example.com" in text
    assert "![картинка](https://img.example.com/i.png)" in text and ", дальше." in text
    assert sorted(p.name for p in out.iterdir()) == ["note_qr.md", "note_qr_001.png", "note_qr_002.png", "note_qr_003.png"]


def test_qrdoc_html(tmp_path):
    import qrdoc
    src = tmp_path / "page.html"
    src.write_text('<p>Смотри <a class="x" href="https://a.example.com/?q=1&amp;r=2">тут</a> и '
                   '<a href="https://b.example.com">там</a>.</p>', encoding="utf-8")
    links = qrdoc.scan(src, tmp_path)["links"]
    assert [(l["url"], l["text"]) for l in links] == [("https://a.example.com/?q=1&r=2", "тут"), ("https://b.example.com", "там")]
    out = tmp_path / "out"; out.mkdir()
    dst, n = qrdoc.apply(src, out, "page", {"https://a.example.com/?q=1&r=2"}, "replace", 3, qrtools.read_style({}), tmp_path / "w")
    html_out = dst.read_text(encoding="utf-8")
    assert n == 1 and html_out.count("data:image/png;base64,") == 1 and ">там</a>" in html_out and ">тут<" not in html_out


def test_qrdoc_endpoints(client):
    import json as _json
    data = make_qr_docx()
    r = client.post("/api/qrdoc/scan", data={"file": (io.BytesIO(data), "Отчёт.docx")}, content_type="multipart/form-data")
    urls = [l["url"] for l in r.get_json()["links"]]
    assert r.status_code == 200 and len(urls) == 3
    r = client.post("/api/qrdoc/apply", data={"file": (io.BytesIO(data), "Отчёт.docx"), "urls": _json.dumps(urls),
                                              "mode": "replace", "size_cm": "2.5"}, content_type="multipart/form-data")
    d = r.get_json()
    assert r.status_code == 200 and d["done"] == 3 and d["single_file"] == "Отчёт_qr.docx"
    got = client.get(f"/api/download/{d['job']}?raw=1")
    assert zipfile.ZipFile(io.BytesIO(got.data)).read("word/document.xml").count(b"<w:drawing>") == 3
    got.close()

    bad = client.post("/api/qrdoc/scan", data={"file": (io.BytesIO(b"x"), "a.pdf")}, content_type="multipart/form-data")
    assert bad.status_code == 400 and "не поддерживается" in bad.get_json()["error"]
    none = client.post("/api/qrdoc/apply", data={"file": (io.BytesIO(data), "a.docx"), "urls": "[]"}, content_type="multipart/form-data")
    assert none.status_code == 400
    broken = client.post("/api/qrdoc/scan", data={"file": (io.BytesIO(b"not zip"), "a.docx")}, content_type="multipart/form-data")
    assert broken.status_code == 400 and "повреждён" in broken.get_json()["error"]


@pytest.mark.skipif(not shutil.which("pandoc"), reason="нет pandoc")
def test_qrdoc_odt_via_docx(tmp_path):
    import qrdoc
    md = tmp_path / "s.md"
    md.write_text("[Ссылка](https://odt.example.com)\n", encoding="utf-8")
    subprocess.run(["pandoc", str(md), "-o", str(tmp_path / "s.odt")], check=True)
    out = tmp_path / "out"; out.mkdir()
    dst, n = qrdoc.apply(tmp_path / "s.odt", out, "s", {"https://odt.example.com"}, "replace", 3, qrtools.read_style({}), tmp_path / "w")
    assert dst.suffix == ".docx" and n == 1


# ---------------------------------------------------------------- загрузчик: отмена и скорость

@pytest.fixture()
def slow_server(media, tmp_path):
    """/big.mp4 отдаётся по ~100 КБ/с, /hang вообще не отвечает (как зависший сайт)."""
    import http.server
    import threading
    big = tmp_path / "big.bin"
    big.write_bytes(os.urandom(4 * 1048576))

    class Slow(http.server.BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_GET(self):
            if self.path.startswith("/hang"):
                time.sleep(60)
                return
            data = big.read_bytes()
            self.send_response(200)
            self.send_header("Content-Type", "video/mp4")
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            try:
                for i in range(0, len(data), 8192):
                    self.wfile.write(data[i:i + 8192])
                    time.sleep(0.08)
            except OSError:
                pass

    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Slow)
    srv.daemon_threads = True
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_port}"
    srv.shutdown()


def _cancel_and_time(client, tid):
    t0 = time.time()
    client.post(f"/api/task/{tid}/cancel")
    d = wait_task(client, tid, timeout=20)
    return d, time.time() - t0


@pytest.mark.parametrize("path, wait_for", [("/big.mp4", "Качаю"), ("/hang/video.mp4", "получаю данные")])
def test_dl_cancel_is_immediate(client, slow_server, path, wait_for):
    r = client.post("/api/dl", data={"urls": slow_server + path, "mode": "video"})
    tid = r.get_json()["task"]
    t0 = time.time()
    while wait_for not in client.get(f"/api/task/{tid}").get_json()["message"] and time.time() - t0 < 20:
        time.sleep(0.2)
    time.sleep(1)
    task = webapp.TASKS.get(tid)
    worker = task.proc
    assert worker is not None and worker.poll() is None            # качает отдельный процесс
    d, took = _cancel_and_time(client, tid)
    assert d["status"] == "cancelled" and took < 3, (d, took)
    assert worker.poll() is not None                              # процесс убит, а не брошен
    assert not any(task.out_dir.iterdir())                        # недокачанные .part убраны


def test_dl_worker_crash_is_reported(client, monkeypatch):
    import downloader
    monkeypatch.setattr(downloader, "worker_cmd", lambda args: [sys.executable, "-c", "import sys; sys.exit('boom')"])
    d = wait_task(client, client.post("/api/dl", data={"urls": "https://example.com/v"}).get_json()["task"])
    assert d["status"] == "error" and "boom" in d["error"]


def test_dl_opts_speed_and_js_runtime():
    y = downloader.build_opts(downloader.read_opts({"mode": "video"}), Path("/tmp"), None)
    assert y["concurrent_fragment_downloads"] == 4 and "deno" in y["js_runtimes"]


def test_engine_deno_install_sets_exec_bit(engine_env, monkeypatch):
    engines, srv_dir, base = engine_env
    exe = "deno.exe" if sys.platform == "win32" else "deno"
    with zipfile.ZipFile(srv_dir / engines._deno_asset(), "w") as z:
        z.writestr(exe, "#!/bin/sh\necho deno 2.0\n")
    monkeypatch.setitem(engines.CATALOG, "deno", {**engines.CATALOG["deno"], "url": base + "/{asset}"})
    path = engines.install("deno")
    assert Path(path, exe).exists() and engines.can_install("deno")
    if sys.platform != "win32":
        assert os.access(Path(path, exe), os.X_OK)


def test_dl_read_events_windows_newlines_and_partial_tail(tmp_path):
    """Файл событий с \\r\\n (как пишет Windows в текстовом режиме) и недописанной последней строкой."""
    p = tmp_path / "ev.jsonl"
    p.write_bytes(b'{"message": "a"}\r\n{"done": "u"}\r\n\r\n{"progress": 0.5, "mess')
    evs, pos = downloader.read_events(p, 0)
    assert evs == [{"message": "a"}, {"done": "u"}]
    with p.open("ab") as f:
        f.write(b'age": "b"}\n')
    evs, pos2 = downloader.read_events(p, pos)
    assert evs == [{"progress": 0.5, "message": "b"}] and pos2 == p.stat().st_size
    assert downloader.read_events(tmp_path / "none", 7) == ([], 7)


# ---------------------------------------------------------------- сеть и VPN

def test_net_settings(monkeypatch):
    import urllib.request
    rn = downloader.read_net
    assert rn({}) == {"proxy": "auto", "ipv4": True}
    assert rn({"proxy": "127.0.0.1:10809", "ipv4": "0"}) == {"proxy": "http://127.0.0.1:10809", "ipv4": False}
    assert rn({"proxy": "socks5://127.0.0.1:10808"})["proxy"] == "socks5://127.0.0.1:10808"
    assert rn({"proxy": "ftp://x; rm -rf"})["proxy"] == "auto"
    assert downloader.net_opts({"proxy": "none", "ipv4": True}) == {"proxy": "", "source_address": "0.0.0.0"}
    assert downloader.net_opts({"proxy": "http://p:1", "ipv4": False}) == {"proxy": "http://p:1"}
    # Windows отдаёт socks из реестра как socks4, а v2rayN/Clash говорят socks5
    monkeypatch.setattr(urllib.request, "getproxies", lambda: {"socks": "socks4://127.0.0.1:10808"})
    assert downloader.system_proxy() == "socks5h://127.0.0.1:10808"
    assert downloader.net_opts({"proxy": "auto", "ipv4": False}) == {"proxy": "socks5h://127.0.0.1:10808"}
    monkeypatch.setattr(urllib.request, "getproxies", lambda: {"https": "http://127.0.0.1:10809"})
    assert downloader.net_opts({"proxy": "auto", "ipv4": False}) == {}   # yt-dlp возьмёт системный сам
    assert [a["proxy"] for a in downloader._net_attempts({"proxy": "auto"})] == ["auto", "none"]
    monkeypatch.setattr(urllib.request, "getproxies", lambda: {})
    assert len(downloader._net_attempts({"proxy": "auto"})) == 1


def test_net_human_errors():
    h = downloader._human_error
    assert "VPN" in h("ERROR: [youtube] x: Sign in to confirm you’re not a bot. Use --cookies")
    assert "прокси" in h("Unable to connect to proxy: [Errno 111] Connection refused")
    assert "Проверить сеть" in h("<urlopen error timed out>")


@needs_ffmpeg
def test_dl_dead_system_proxy_falls_back_to_direct(client, http_media, monkeypatch):
    """VPN-клиент выключили, а прокси в системе остался: качаем напрямую."""
    dead = "http://127.0.0.1:9"                       # discard-порт, там никого нет
    for k in ("http_proxy", "HTTP_PROXY", "https_proxy", "HTTPS_PROXY", "all_proxy", "ALL_PROXY"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("NO_PROXY", "")
    monkeypatch.setenv("no_proxy", "")
    monkeypatch.setenv("http_proxy", dead)
    monkeypatch.setenv("https_proxy", dead)
    d = client.post("/api/dl/info", data={"urls": f"{http_media}/a.mp4"}).get_json()
    assert d["items"][0]["title"] == "a", d
    r = client.post("/api/dl", data={"urls": f"{http_media}/a.mp4", "mode": "video"})
    d = wait_task(client, r.get_json()["task"])
    assert d["status"] == "done" and d["single_file"].endswith(".mp4"), d
    # а если прокси задан явно и он мёртвый — понятная ошибка, без тихого обхода
    r = client.post("/api/dl", data={"urls": f"{http_media}/a.mp4", "proxy": dead})
    d = wait_task(client, r.get_json()["task"])
    assert d["status"] == "error" and "прокси" in d["error"], d


def test_dl_netcheck(client, monkeypatch):
    calls = []

    class FakeYDL:
        def __init__(self, opts):
            calls.append(opts)

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def urlopen(self, url):
            if "vk.com" in url:
                raise OSError("<urlopen error timed out>")
            return io.BytesIO(b"x")

    monkeypatch.setattr(downloader.yt_dlp, "YoutubeDL", FakeYDL)
    d = client.post("/api/dl/netcheck", data={"proxy": "none", "ipv4": "1"}).get_json()
    by = {x["name"]: x for x in d["sites"]}
    assert by["YouTube"]["ok"] and not by["VK"]["ok"] and by["VK"]["error"] == "нет ответа"
    assert downloader._short_net_error("<urlopen error Tunnel connection failed: 403 Forbidden>") == "прокси не пустил (403)"
    assert d["proxy"] == "" and d["ipv4"] is True and "ipv6" in d
    assert all(c["proxy"] == "" and c["source_address"] == "0.0.0.0" for c in calls)




# ---------------------------------------------------------------- сервер для телефона

@pytest.fixture()
def phone(tmp_path, monkeypatch):
    import phoneserver
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path))
    monkeypatch.setattr(phoneserver, "PORTS", range(18700, 18720))
    yield phoneserver
    phoneserver.stop(remember=False)


def _get(url, password=None):
    import base64
    import urllib.error
    import urllib.request
    req = urllib.request.Request(url)
    if password is not None:
        req.add_header("Authorization", "Basic " + base64.b64encode(f"телефон:{password}".encode()).decode())
    try:
        with urllib.request.urlopen(req, timeout=10) as r:
            return r.status, r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def test_phone_server_lifecycle(client, phone):
    d = client.get("/api/phone").get_json()
    assert d["running"] is False and d["enabled"] is False
    d = client.post("/api/phone", data={"action": "on"}).get_json()
    assert d["running"] and d["enabled"] and len(d["password"]) == 8 and d["port"] in phone.PORTS
    base = f"http://127.0.0.1:{d['port']}"
    # через сервер для телефона без пароля и с чужим не пускает, с правильным пускает (логин любой)
    assert _get(base + "/api/formats")[0] == 401
    assert _get(base + "/api/formats", "wrong")[0] == 401
    st, body = _get(base + "/api/formats", d["password"])
    assert st == 200 and b"categories" in body
    # с телефона сервер не выключить и пароль не подсмотреть
    assert _get(base + "/api/phone", d["password"])[0] == 403
    if d["urls"]:
        assert d["link"].startswith("universal://connect?url=http%3A%2F%2F") and d["password"] in d["link"]
        assert client.get("/api/phone/qr.svg").data.startswith(b"<svg")
    # новый пароль сразу заменяет старый
    old = d["password"]
    d = client.post("/api/phone", data={"action": "password"}).get_json()
    assert d["password"] != old
    assert _get(base + "/api/formats", old)[0] == 401
    assert _get(base + "/api/formats", d["password"])[0] == 200
    # настройка запоминается: после перезапуска программы сервер включится сам на том же порту
    port = d["port"]
    phone.stop(remember=False)
    phone.autostart(webapp.app)
    assert phone.running() and phone.settings()["port"] == port
    d = client.post("/api/phone", data={"action": "off"}).get_json()
    assert not d["running"] and not d["enabled"]
    import urllib.error
    with pytest.raises(urllib.error.URLError):                     # порт закрыт
        _get(base + "/api/formats", "x")


def test_phone_api_only_from_this_computer(client, phone):
    assert client.get("/api/phone", environ_base={"REMOTE_ADDR": "192.168.1.5"}).status_code == 403
    assert client.post("/api/phone", data={"action": "on"}, environ_base={"REMOTE_ADDR": "192.168.1.5"}).status_code == 403
    assert not phone.running()


def test_phone_link_and_addresses(phone):
    link = phone.link("http://192.168.1.20:5000", "ab cd")
    assert link == "universal://connect?url=http%3A%2F%2F192.168.1.20%3A5000&password=ab%20cd"
    assert all(not ip.startswith("127.") for ip in phone.addresses())
    assert set(phone.new_password()) <= set(phone.ALPHABET)


# ---------------------------------------------------------------- распознавание текста (OCRmyPDF)

import ocr  # noqa: E402

needs_ocr = pytest.mark.skipif(not ocr.available() or "eng" not in ocr.languages(), reason="нет OCRmyPDF или Tesseract")


def _scan(path: Path, lines: list[str], angle: float = 0) -> None:
    from PIL import Image, ImageDraw, ImageFont
    im = Image.new("RGB", (1654, 1169), "white")
    d = ImageDraw.Draw(im)
    font = None
    for f in ("DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "arial.ttf",
              "/System/Library/Fonts/Supplemental/Arial.ttf"):
        try:
            font = ImageFont.truetype(f, 56)
            break
        except OSError:
            continue
    for i, line in enumerate(lines):
        d.text((120, 200 + 110 * i), line, fill="black", font=font or ImageFont.load_default())
    if angle:
        im = im.rotate(angle, expand=True, fillcolor="white")
    if path.suffix == ".pdf":
        im.save(path, resolution=200)
    else:
        im.save(path, dpi=(200, 200))


def test_ocr_args():
    o = ocr.read_opts({"langs": "rus+eng", "mode": "redo", "deskew": "1", "rotate": "1", "pdfa": "1"})
    a = ocr.build_args(Path("in.pdf"), Path("out.pdf"), o, None, True)
    assert "--redo-ocr" in a and "--deskew" not in a and "--rotate-pages" not in a   # redo с ними несовместим
    assert a[a.index("-l") + 1] == "rus+eng" and a[-2:] == ["in.pdf", "out.pdf"]
    o = ocr.read_opts({"langs": "eng", "mode": "skip"})
    a = ocr.build_args(Path("photo.jpg"), Path("o.pdf"), o, Path("o.txt"), False)
    assert "--skip-text" in a and "--deskew" in a and "--image-dpi" in a and "--sidecar" in a
    assert a[a.index("--output-type") + 1] == "pdf"
    with pytest.raises(ocr.OcrError):
        ocr.read_opts({"langs": "rus;rm -rf"})


def test_ocr_pdfa_without_ghostscript(monkeypatch):
    monkeypatch.setattr(ocr, "_ghostscript", lambda: None)
    monkeypatch.setattr(ocr, "version", lambda: "17.13.0")
    a = ocr.build_args(Path("a.pdf"), Path("b.pdf"), ocr.read_opts({"langs": "eng"}), None, True)
    assert a[a.index("--pdfa-backend") + 1] == "internal"
    # старый OCRmyPDF такого ключа не знает и без Ghostscript PDF/A не сделает: сохраняем обычный PDF
    monkeypatch.setattr(ocr, "version", lambda: "17.12.1")
    a = ocr.build_args(Path("a.pdf"), Path("b.pdf"), ocr.read_opts({"langs": "eng"}), None, True)
    assert "--pdfa-backend" not in a and a[a.index("--output-type") + 1] == "pdf"
    monkeypatch.setattr(ocr, "_ghostscript", lambda: "gs")
    a = ocr.build_args(Path("a.pdf"), Path("b.pdf"), ocr.read_opts({"langs": "eng"}), None, True)
    assert "--pdfa-backend" not in a and a[a.index("--output-type") + 1] == "pdfa"


@needs_ocr
def test_ocr_scan_and_photo(client, tmp_path):
    from pypdf import PdfReader
    _scan(tmp_path / "scan.pdf", ["Hello OCR test 12345", "Universal recognises text"])
    _scan(tmp_path / "photo.png", ["Skewed photo of a page", "Second line 67890"], angle=4)
    files = [(io.BytesIO((tmp_path / n).read_bytes()), n) for n in ("scan.pdf", "photo.png")]
    r = client.post("/api/ocr", data={"files": files, "langs": "eng", "txt": "1"}, content_type="multipart/form-data")
    d = wait_task(client, r.get_json()["task"], timeout=300)
    assert d["status"] == "done", d
    assert sorted(d["names"]) == ["photo.txt", "photo_ocr.pdf", "scan.txt", "scan_ocr.pdf"]
    out = webapp.WORK_ROOT / d["job"] / "out"
    text = PdfReader(out / "scan_ocr.pdf").pages[0].extract_text()
    assert "12345" in text and "Universal" in text
    assert "67890" in (out / "photo.txt").read_text(encoding="utf-8")


@needs_ocr
def test_ocr_pdfa_and_existing_text(client, tmp_path):
    from pypdf import PdfReader
    _scan(tmp_path / "scan.pdf", ["Archive copy 2468"])
    r = client.post("/api/ocr", data={"files": [(io.BytesIO((tmp_path / "scan.pdf").read_bytes()), "scan.pdf")],
                                      "langs": "eng", "pdfa": "1"}, content_type="multipart/form-data")
    d = wait_task(client, r.get_json()["task"], timeout=300)
    assert d["status"] == "done", d
    done = webapp.WORK_ROOT / d["job"] / "out" / "scan_ocr.pdf"
    assert b"pdfaid" in done.read_bytes() or d["failed"]           # PDF/A, если сошёлся
    # PDF уже с текстом: «пропустить» не ломается, текст остаётся
    r = client.post("/api/ocr", data={"files": [(io.BytesIO(done.read_bytes()), "again.pdf")], "langs": "eng"},
                    content_type="multipart/form-data")
    d = wait_task(client, r.get_json()["task"], timeout=300)
    assert d["status"] == "done", d
    assert "2468" in PdfReader(webapp.WORK_ROOT / d["job"] / "out" / "again_ocr.pdf").pages[0].extract_text()


@needs_ocr
def test_ocr_unknown_language(client, tmp_path):
    _scan(tmp_path / "scan.pdf", ["x"])
    r = client.post("/api/ocr", data={"files": [(io.BytesIO((tmp_path / "scan.pdf").read_bytes()), "scan.pdf")],
                                      "langs": "klingon"}, content_type="multipart/form-data")
    d = wait_task(client, r.get_json()["task"])
    assert d["status"] == "error" and "нет языков" in d["error"]


def test_engine_index_and_any_release(monkeypatch):
    import engines
    page = ('<a href="tesseract-ocr-w64-setup-5.3.0.20221222.exe">a</a>'
            '<a href="tesseract-ocr-w64-setup-5.5.0.20241111.exe">b</a><a href="other.zip">c</a>')

    class R(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    monkeypatch.setattr(engines, "_get", lambda url, timeout=30: R(page.encode()))
    assets = engines._assets_from_index("https://example.test/tesseract/")
    assert assets[0] == ("tesseract-ocr-w64-setup-5.5.0.20241111.exe",
                         "https://example.test/tesseract/tesseract-ocr-w64-setup-5.5.0.20241111.exe")
    # у последнего релиза установщика нет: берётся из более раннего, а не ошибка
    monkeypatch.setattr(engines, "_assets_from_api", lambda repo: [("tesseract-5.5.2.tar.gz", "u0")])
    monkeypatch.setattr(engines, "_assets_from_releases", lambda repo: [
        ("tesseract-5.5.2.tar.gz", "u0"), ("tesseract-ocr-w64-setup-5.5.1.exe", "u1")])
    assert engines._resolve_url(engines.CATALOG["tesseract"]) == "u1"


# ---------------------------------------------------------------- обновления

import updater  # noqa: E402

FAKE_RELEASE = {"version": "99.1.0", "notes": "- **Новое**: что-то\n- ещё [ссылка](https://x)\n", "assets": [
    ("Universal-Setup-99.1.0.exe", "https://example.test/setup.exe", 30_000_000),
    ("Universal-99.1.0-macos-arm64.dmg", "https://example.test/arm.dmg", 1),
    ("Universal-99.1.0-macos-x86_64.dmg", "https://example.test/intel.dmg", 1),
    ("universal_99.1.0_all.deb", "https://example.test/u.deb", 1),
    ("universal-99.1.0-1.noarch.rpm", "https://example.test/u.rpm", 1),
    ("universal-99.1.0-1-any.pkg.tar.zst", "https://example.test/u.pkg.tar.zst", 1),
    ("Universal-99.1.0-android.apk", "https://example.test/u.apk", 1)]}


@pytest.fixture()
def fake_release(tmp_path, monkeypatch):
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path))
    calls = []
    monkeypatch.setattr(updater, "_fetch", lambda: calls.append(1) or json.loads(json.dumps(FAKE_RELEASE)))
    return calls


def test_update_check_and_cache(fake_release, monkeypatch):
    monkeypatch.setattr(updater, "kind", lambda: "installer")
    d = updater.check()
    assert d["newer"] and d["latest"] == "99.1.0" and d["current"] == updater.VERSION
    assert d["asset"]["name"] == "Universal-Setup-99.1.0.exe" and d["can_install"]
    assert d["url"].endswith("/releases/tag/v99.1.0")
    updater.check()
    assert len(fake_release) == 1                         # второй раз из кэша, GitHub не дёргаем
    updater.check(force=True)
    assert len(fake_release) == 2


def test_update_assets_per_platform(fake_release, monkeypatch):
    want = {"deb": "universal_99.1.0_all.deb", "rpm": "universal-99.1.0-1.noarch.rpm",
            "arch": "universal-99.1.0-1-any.pkg.tar.zst"}
    for k, name in want.items():
        monkeypatch.setattr(updater, "kind", lambda k=k: k)
        d = updater.check()
        assert d["asset"]["name"] == name and d["can_install"]
    monkeypatch.setattr(updater, "kind", lambda: "source")     # из исходников: только ссылка, обновление через git pull
    assert updater.check()["asset"] is None and not updater.check()["can_install"]
    monkeypatch.setattr(updater, "kind", lambda: "dmg")
    monkeypatch.setattr(updater.platform, "machine", lambda: "x86_64")
    assert updater.check()["asset"]["name"] == "Universal-99.1.0-macos-x86_64.dmg"
    monkeypatch.setattr(updater.platform, "machine", lambda: "arm64")
    assert updater.check()["asset"]["name"] == "Universal-99.1.0-macos-arm64.dmg"
    monkeypatch.setattr(updater, "kind", lambda: "docker")
    assert not updater.check()["can_install"]


def test_update_not_newer(fake_release, monkeypatch):
    monkeypatch.setattr(updater, "VERSION", "99.1.0")
    assert not updater.check()["newer"]
    monkeypatch.setattr(updater, "VERSION", "100.0.0")
    assert not updater.check()["newer"]
    assert updater.parse("1.10.0") > updater.parse("1.9.9")


def test_update_install_windows_and_linux(fake_release, monkeypatch, client):
    import engines
    monkeypatch.setattr(engines, "_download", lambda url, dest, task, a, b: Path(dest).write_bytes(b"x"))
    started, opened, quit_calls = [], [], []
    monkeypatch.setattr(updater.subprocess, "Popen", lambda args, **kw: started.append(args))
    monkeypatch.setattr(updater, "_open", lambda p: opened.append(p))
    monkeypatch.setattr(updater.threading, "Timer", lambda s, fn: type("T", (), {"start": lambda self: quit_calls.append(fn)})())
    monkeypatch.setattr(updater, "kind", lambda: "installer")
    webapp.app.config["QUIT_APP"] = "quit"
    try:
        d = wait_task(client, client.post("/api/update").get_json()["task"])
    finally:
        webapp.app.config.pop("QUIT_APP")
    assert d["status"] == "done" and d["result"]["action"] == "restart"
    assert started[0][0].endswith("Universal-Setup-99.1.0.exe") and "/SILENT" in started[0]
    assert quit_calls == ["quit"]                          # приложение закроется, чтобы установщик его заменил
    monkeypatch.setattr(updater, "kind", lambda: "deb")
    d = wait_task(client, client.post("/api/update").get_json()["task"])
    assert d["status"] == "done" and d["result"]["command"].startswith("sudo apt install ")
    assert opened and str(opened[0]).endswith("universal_99.1.0_all.deb")
    # прошлое скачанное обновление удалено
    assert sorted(p.name for p in (updater.engines.root() / "updates").iterdir()) == ["universal_99.1.0_all.deb"]


def test_update_api_local_only(fake_release, client):
    assert client.get("/api/update", environ_base={"REMOTE_ADDR": "192.168.1.5"}).status_code == 403
    assert client.post("/api/update", environ_base={"REMOTE_ADDR": "192.168.1.5"}).status_code == 403
    assert client.get("/api/update").get_json()["latest"] == "99.1.0"


def test_update_github_unreachable(tmp_path, monkeypatch, client):
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path))

    def boom():
        raise updater.UpdateError("не удалось связаться с GitHub (URLError)")
    monkeypatch.setattr(updater, "_fetch", boom)
    r = client.get("/api/update")
    assert r.status_code == 502 and "GitHub" in r.get_json()["error"] and r.get_json()["current"] == updater.VERSION


# ---------------------------------------------------------------- информация

def test_about_info(client, phone):
    d = client.get("/api/about").get_json()
    assert d["version"] == updater.VERSION and d["home"] and d["python"]
    names = {t["name"] for t in d["tools"]}
    assert {"ffmpeg", "yt-dlp", "OCRmyPDF", "Tesseract"} <= names
    ff = next(t for t in d["tools"] if t["name"] == "ffmpeg")
    if shutil.which("ffmpeg"):
        assert ff["installed"] and ff["version"] and ff["path"]
    # с телефона (через сервер для телефона) пути на компьютере не отдаются
    on = client.post("/api/phone", data={"action": "on"}).get_json()
    st, body = _get(f"http://127.0.0.1:{on['port']}/api/about", on["password"])
    remote = json.loads(body)
    assert st == 200 and "home" not in remote and "jobs" not in remote and all(not t["path"] for t in remote["tools"])
    assert client.post("/api/about/open", environ_base={"REMOTE_ADDR": "10.0.0.2"}).status_code == 403


def test_about_nerd(client):
    d = client.get("/api/about/nerd").get_json()
    titles = [s["title"] for s in d["sections"]]
    assert titles[:2] == ["Приложение", "Система"] and "Python-пакеты" in titles and "Движки" in titles
    app_rows = dict(d["sections"][0]["rows"])
    assert app_rows["Версия"] == updater.VERSION and app_rows["PID"] == str(os.getpid())
    pk = dict(next(s for s in d["sections"] if s["title"] == "Python-пакеты")["rows"])
    assert pk["flask"] != "нет"
    assert client.get("/api/about/nerd", environ_base={"REMOTE_ADDR": "192.168.1.9"}).status_code == 403


# ---------------------------------------------------------------- архивы и файлы

def _files_post(client, files, **form):
    data = {"files": [(io.BytesIO(b), n) for n, b in files], "relpath": [r for r, _ in files], **form}
    r = client.post("/api/files", data=data, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    return wait_task(client, r.get_json()["task"])


def test_files_zip_keeps_folders_and_downloads_nested(client):
    d = _files_post(client, [("папка/a.txt", b"one"), ("папка/вложенная/b.txt", b"two")], op="zip", name="мой архив")
    assert d["status"] == "done" and d["names"] == ["мой архив.zip"]
    z = zipfile.ZipFile(webapp.WORK_ROOT / d["job"] / "out" / "мой архив.zip")
    assert sorted(z.namelist()) == ["папка/a.txt", "папка/вложенная/b.txt"] and z.read("папка/вложенная/b.txt") == b"two"
    # распаковка того же архива: папки сохраняются, отдельный файл качается по вложенному пути
    zbytes = (webapp.WORK_ROOT / d["job"] / "out" / "мой архив.zip").read_bytes()
    d = _files_post(client, [("мой архив.zip", zbytes)], op="extract")
    assert d["status"] == "done" and sorted(d["names"]) == ["мой архив/папка/a.txt", "мой архив/папка/вложенная/b.txt"]
    assert client.get(f"/api/file/{d['job']}/мой архив/папка/вложенная/b.txt").data == b"two"
    assert client.get(f"/api/file/{d['job']}/../../etc/passwd").status_code == 404
    bundle = zipfile.ZipFile(io.BytesIO(client.get(f"/api/download/{d['job']}").data))
    assert "мой архив/папка/вложенная/b.txt" in bundle.namelist()


def test_files_extract_windows_names_and_zip_slip(client, monkeypatch):
    buf = io.BytesIO()
    # «Проводник» Windows пишет имена в cp866 без флага UTF-8; zipfile Python так сам не умеет — подменяем кодирование
    monkeypatch.setattr(zipfile.ZipInfo, "_encodeFilenameFlags",
                        lambda self: (self.filename.encode("cp866", "replace"), self.flag_bits & ~0x800))
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("Документ.txt", b"x")
        z.writestr("../../evil.txt", b"bad")
    monkeypatch.undo()
    d = _files_post(client, [("win.zip", buf.getvalue())], op="extract")
    assert d["status"] == "done"
    assert "win/Документ.txt" in d["names"] and all("evil" not in n or n.startswith("win/") for n in d["names"])
    assert not (webapp.WORK_ROOT / d["job"] / "evil.txt").exists()


@pytest.mark.skipif(not __import__("filetools").has_7z(), reason="нет py7zr")
def test_files_7z_with_password(client):
    d = _files_post(client, [("секрет.txt", "тайна".encode())], op="7z", name="s", password="пароль")
    assert d["status"] == "done" and d["names"] == ["s.7z"]
    data = (webapp.WORK_ROOT / d["job"] / "out" / "s.7z").read_bytes()
    bad = _files_post(client, [("s.7z", data)], op="extract", password="не тот")
    assert bad["status"] == "error"
    ok = _files_post(client, [("s.7z", data)], op="extract", password="пароль")
    assert ok["status"] == "done" and ok["names"] == ["s/секрет.txt"]


def test_files_hash_and_rename(client):
    import hashlib
    d = _files_post(client, [("a.bin", b"abc"), ("b.bin", b"")], op="hash", algos="md5,sha256")
    rows = {r["name"]: r for r in d["result"]["hashes"]}
    assert rows["a.bin"]["sha256"] == hashlib.sha256(b"abc").hexdigest() and rows["b.bin"]["md5"] == hashlib.md5(b"").hexdigest()
    assert sorted(d["names"]) == ["MD5SUMS.txt", "SHA256SUMS.txt"]
    sums = (webapp.WORK_ROOT / d["job"] / "out" / "SHA256SUMS.txt").read_text(encoding="utf-8")
    assert f"{hashlib.sha256(b'abc').hexdigest()}  a.bin" in sums
    d = _files_post(client, [("IMG_1.JPG", b"1"), ("IMG_2.JPG", b"2")], op="rename", template="отпуск_{n}", start="5",
                    digits="3", case="lower")
    assert sorted(d["names"]) == ["отпуск_005.jpg", "отпуск_006.jpg"]
    r = client.post("/api/files", data={"files": [(io.BytesIO(b"x"), "a")], "op": "rename", "template": "same"},
                    content_type="multipart/form-data")
    assert r.status_code == 400                               # без {name}/{n} все получили бы одно имя


# ---------------------------------------------------------------- редактор документов и фоторедактор

def test_doc_open_save_text_and_html(client):
    r = client.post("/api/doc/open", data={"file": (io.BytesIO("строка 1\nстрока <2>".encode()), "заметка.txt")},
                    content_type="multipart/form-data").get_json()
    assert r["name"] == "заметка" and "строка &lt;2&gt;" in r["html"]
    r = client.post("/api/doc/open", data={"file": (io.BytesIO(b"<html><body><p onclick='x()'>hi</p><script>bad()</script></body></html>"), "p.html")},
                    content_type="multipart/form-data").get_json()
    assert "hi" in r["html"] and "script" not in r["html"] and "onclick" not in r["html"]
    s = client.post("/api/doc/save", data={"html": "<h1>Тест</h1><p>текст</p>", "format": "html", "name": "отчёт"}).get_json()
    body = client.get(s["url"]).data.decode()
    assert s["name"] == "отчёт.html" and "<h1>Тест</h1>" in body


@pytest.mark.skipif(not shutil.which("pandoc"), reason="нет pandoc")
def test_doc_roundtrip_docx(client):
    s = client.post("/api/doc/save", data={"html": "<h1>Заголовок</h1><p><b>жирный</b> текст</p><table><tr><td>A1</td></tr></table>",
                                           "format": "docx", "name": "doc"}).get_json()
    data = client.get(s["url"]).data
    assert data[:2] == b"PK"
    r = client.post("/api/doc/open", data={"file": (io.BytesIO(data), "doc.docx")}, content_type="multipart/form-data").get_json()
    assert "Заголовок" in r["html"] and "жирный" in r["html"] and "<table" in r["html"].lower()


def test_photo_load_converts_for_browser(client):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (30, 20), (10, 200, 30)).save(buf, "TIFF")
    r = client.post("/api/photo/load", data={"file": (io.BytesIO(buf.getvalue()), "scan.tif")}, content_type="multipart/form-data")
    assert r.status_code == 200 and r.data[:8] == b"\x89PNG\r\n\x1a\n"
    assert Image.open(io.BytesIO(r.data)).size == (30, 20)


def test_static_editors_served(client):
    for path in ("/static/docedit.js", "/static/photoedit.js", "/static/vendor/jodit/jodit.min.js"):
        assert client.get(path).status_code == 200


def test_pdf_editor_assets(client):
    # pdf.js грузится как ES-модуль: без JS-типа WebView его не выполнит
    for path in ("/static/pdfedit.js", "/static/vendor/pdfjs/pdf.min.js", "/static/vendor/pdfjs/pdf.worker.min.js",
                 "/static/vendor/pdflib/pdf-lib.min.js"):
        r = client.get(path)
        assert r.status_code == 200 and "javascript" in r.content_type
    assert client.get("/static/vendor/pdfjs/cmaps/78-H.bcmap").status_code == 200
    page = client.get("/").get_data(as_text=True)
    assert 'id="pdfView"' in page and "pdfedit.js" in page and "'Редактор PDF'" in page


def test_changelog_since():
    import changelog
    text = "# История\n\n## v2.1.0 — 2026-11-01\n- новое\n\n## v2.0.0 — 2026-10-20\n### Раздел\n- ещё\n\n## v1.9.0 — 2026-10-01\n- старое\n"
    assert [s["version"] for s in changelog.since("1.9.0", "2.1.0", text=text)] == ["2.1.0", "2.0.0"]
    assert [s["version"] for s in changelog.since("", "2.0.0", text=text)] == ["2.0.0"]
    assert changelog.since("2.0.0", "2.0.0", text=text) == []
    assert changelog.sections(text)[1]["md"].startswith("### Раздел")


def test_whatsnew_route(client):
    import updater
    d = client.get("/api/whatsnew").get_json()
    assert d["version"] == updater.VERSION
    assert d["sections"] and d["sections"][0]["version"] == updater.VERSION   # CHANGELOG.md не отстаёт от version.py
    assert client.get("/api/whatsnew?since=" + updater.VERSION).get_json()["sections"] == []


def test_whatsnew_after_update_once(tmp_path, monkeypatch):
    import changelog
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path))
    monkeypatch.setattr(changelog, "sections", lambda text=None: [
        {"version": "2.1.0", "date": "", "md": "- b"}, {"version": "2.0.0", "date": "", "md": "- a"}])
    (tmp_path / "seen_version.txt").write_text("2.0.0", encoding="utf-8")
    assert [s["version"] for s in changelog.after_update("2.1.0")] == ["2.1.0"]
    assert changelog.after_update("2.1.0") == []          # второй запуск — окна нет


@pytest.mark.skipif(not shutil.which("ffmpeg"), reason="нужен ffmpeg")
def test_media_progress_reported(tmp_path):
    import converters as conv
    src = tmp_path / "tone.wav"
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "sine=frequency=440:duration=20", str(src)], check=True)
    seen = []
    out = conv.convert_media(src, tmp_path, "tone", "mp3", progress=seen.append)
    assert out.exists() and seen and max(seen) > 0.5 and all(0 <= x <= 1 for x in seen)


def test_heic_via_pillow(tmp_path):
    pytest.importorskip("pillow_heif")
    from PIL import Image
    import converters as conv
    src = tmp_path / "a.png"
    Image.new("RGB", (40, 30), "red").save(src)
    out = tmp_path / "a.heic"
    assert conv._pillow_encode(src, out, "heic", 70, True, 0)
    import pillow_heif
    pillow_heif.register_heif_opener()
    assert Image.open(out).size == (40, 30)


def test_ffprobe_utf8_output(monkeypatch):
    # в Windows text=True без encoding читал UTF-8 вывод как cp1252, поток падал и probe видел пустой файл
    import inspect
    import converters as conv
    assert 'encoding="utf-8"' in inspect.getsource(editor.probe)
    assert 'encoding="utf-8"' in inspect.getsource(conv._run)


# ---------- голос в текст ----------

def test_stt_task_writes_txt_and_srt(tmp_path, monkeypatch):
    import stt
    from tasks import Task
    monkeypatch.setattr(stt, "available", lambda: True)
    monkeypatch.setattr(stt, "ensure_model", lambda key, task=None: tmp_path)
    monkeypatch.setattr(stt, "_model", lambda key: None)
    monkeypatch.setattr(stt, "transcribe", lambda src, key, lang, task=None, base=0, span=1, audio=None:
                        ([(0.0, 1.5, "Привет"), (1.5, 3.25, "это проверка")], "ru", 3.25))
    job = tmp_path / "job"
    (job / "out").mkdir(parents=True)
    src = tmp_path / "voice.ogg"
    src.write_bytes(b"x")
    task = Task("abc", "stt", job, "text")
    stt.run_stt_task(task, [(src, "voice")], stt.read_opts({"model": "fast", "lang": "ru", "srt": "1"}))
    assert (job / "out" / "voice.txt").read_text(encoding="utf-8").strip() == "Привет это проверка"
    srt = (job / "out" / "voice.srt").read_text(encoding="utf-8")
    assert "00:00:01,500 --> 00:00:03,250" in srt and srt.startswith("1\n")
    assert task.result["transcripts"][0] == {"name": "voice.ogg", "text": "Привет это проверка", "language": "Русский", "duration": 3.2}


def test_stt_model_download_with_mirror_fallback(tmp_path, monkeypatch):
    import http.server
    import threading
    import stt
    files = {"config.json": b"{}", "model.bin": b"M" * 3000, "tokenizer.json": b"{}"}

    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            if self.path.startswith("/api/models/"):
                body = json.dumps({"siblings": [{"rfilename": k, "size": len(v)} for k, v in files.items()] +
                                  [{"rfilename": ".gitattributes"}, {"rfilename": "README.md"}]}).encode()
            else:
                body = files.get(self.path.rsplit("/", 1)[-1])
            if body is None:
                self.send_response(404); self.end_headers(); return
            self.send_response(200); self.send_header("Content-Length", str(len(body))); self.end_headers(); self.wfile.write(body)

        def log_message(self, *a):
            pass

    srv = http.server.HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path))
    monkeypatch.setattr(stt, "ENDPOINTS", ("http://127.0.0.1:1", f"http://127.0.0.1:{srv.server_port}"))   # первый недоступен
    try:
        d = stt.ensure_model("fast")
    finally:
        srv.shutdown()
    assert stt.downloaded("fast") and (d / "model.bin").stat().st_size == 3000
    assert not (d / "README.md").exists()


def test_stt_route_needs_files(client, monkeypatch):
    import stt
    monkeypatch.setattr(stt, "available", lambda: True)
    assert client.post("/api/stt", data={"model": "fast"}).status_code == 400
    assert client.post("/api/stt", data={"model": "huge"}).get_json()["error"] == "Неизвестная модель"
    assert "stt" in client.get("/api/formats").get_json()["engines"]


@pytest.mark.skipif(not os.environ.get("STT_REAL"), reason="настоящая модель: только в CI (STT_REAL=1)")
def test_stt_real_voice_message(tmp_path):
    # голосовое как в Telegram: синтезированная речь (espeak-ng) -> OGG Opus 32 кбит/с, 48 кГц, моно.
    # Голос espeak роботизированный, поэтому слова проверяем на английском (его он произносит чётко),
    # а на русском — что пришёл русский текст. Живая речь распознаётся заметно лучше.
    import re
    import stt

    def voice(name, lang, phrase):
        wav, ogg = tmp_path / f"{name}.wav", tmp_path / f"{name}.ogg"
        subprocess.run(["espeak-ng", "-v", lang, "-s", "130", "-w", str(wav), phrase], check=True)
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-i", str(wav), "-c:a", "libopus", "-b:a", "32k",
                        "-ar", "48000", "-ac", "1", str(ogg)], check=True)
        return ogg

    stt.ensure_model("fast")
    segs, lang, dur = stt.transcribe(voice("en", "en", "Hello. This is a test of speech recognition."), "fast", "")
    text = " ".join(t for _, _, t in segs).lower()
    print("en:", lang, text)
    assert lang == "en" and dur > 1 and sum(w in text for w in ("hello", "test", "speech", "recognition")) >= 2

    segs, lang, dur = stt.transcribe(voice("ru", "ru", "Привет. Это проверка распознавания речи."), "fast", "ru")
    text = " ".join(t for _, _, t in segs)
    print("ru:", text)
    assert len(re.findall(r"[а-яё]", text.lower())) >= 10


# ---------- первый запуск: движки ставятся сами ----------

def test_auto_setup_installs_missing_and_collects_errors(monkeypatch, tmp_path):
    import engines
    from tasks import Task
    monkeypatch.setattr(engines, "setup_plan", lambda: ["ffmpeg", "pandoc", "poppler"])
    monkeypatch.setattr(engines, "can_install", lambda name="": True)
    monkeypatch.setattr(engines, "missing", lambda: ["poppler"])
    done = []

    def fake_install(name, task=None):
        task.update(0.5, "Скачиваю")
        if name == "poppler":
            raise engines.EngineError("нет интернета")
        done.append(name)
    monkeypatch.setattr(engines, "install", fake_install)
    (tmp_path / "out").mkdir()
    t = Task("s", "setup", tmp_path, "setup")
    engines.auto_setup(t)
    assert done == ["ffmpeg", "pandoc"] and t.done == 2
    assert t.failed == [{"name": "poppler", "error": "нет интернета"}] and t.result["missing"] == ["poppler"]


def test_engine_spec_for_macos(monkeypatch):
    import engines
    monkeypatch.setattr(engines.sys, "platform", "darwin")
    ff, pd = engines._spec("ffmpeg"), engines._spec("pandoc")
    assert "url" not in ff and len(ff["urls"]) == 2 and pd["asset"] == r"{arch}-macOS\.zip$"
    assert engines._exe_names("ffmpeg") == ["ffmpeg", "ffprobe"]
    assert engines.can_install("ffmpeg") and engines.can_install("pandoc") and not engines.can_install("poppler")


def test_setup_endpoint(client, monkeypatch):
    import engines
    assert client.get("/api/setup").get_json() == {"task": None, "missing": []}    # не настольная версия
    monkeypatch.setenv("UNIVERSAL_DESKTOP", "1")
    monkeypatch.setattr(engines, "setup_plan", lambda: [])
    d = client.post("/api/setup").get_json()
    assert d["task"] is None and isinstance(d["missing"], list)
    assert client.get("/api/ping").status_code == 204 and webapp.app.config["LAST_PING"] > 0


# ---------- телефон: встроенный Python без ImageMagick, pandoc и poppler ----------

@pytest.fixture
def android(monkeypatch):
    import converters
    monkeypatch.setenv("UNIVERSAL_ANDROID", "1")
    real = converters._which
    monkeypatch.setattr(converters, "_which", lambda *n: None if set(n) & {"magick", "convert", "pandoc", "soffice", "libreoffice"} else real(*n))
    return converters


def test_android_images_via_pillow(android, tmp_path):
    from PIL import Image
    src = tmp_path / "a.png"
    Image.new("RGBA", (64, 40), (255, 0, 0, 128)).save(src)
    for to in ("jpg", "webp", "gif", "bmp", "ico", "tiff", "pdf", "pgm"):
        out = android.convert_image(src, tmp_path, f"a_{to}", to, quality=80)
        assert out.exists() and out.stat().st_size > 0, to
    assert Image.open(tmp_path / "a_jpg.jpg").mode == "RGB"
    assert android.engines()["imagemagick"] and android.engines()["builtin"]
    outs = android.targets_for("image")["native"]
    assert "png" in outs and "psd" not in outs and "eps" not in outs
    with pytest.raises(android.ConvertError):
        android.convert_image(src, tmp_path, "x", "psd")


def test_android_documents_builtin(android, tmp_path):
    src = tmp_path / "t.md"
    src.write_text("# Заголовок\n\nТекст **жирный**\n\n- раз\n- два\n", encoding="utf-8")
    docx = android.convert_document(src, tmp_path, "t", "docx")
    back = android.convert_document(docx, tmp_path, "back", "md")
    text = back.read_text(encoding="utf-8")
    assert "# Заголовок" in text and "**жирный**" in text and "- раз" in text
    assert set(android.targets_for("doc")["native"]) <= {"html", "md", "txt", "docx"}
    import pydocs
    pdf = pydocs.write("<h1>Привет</h1><p>мир <code>код</code></p>", "pdf", tmp_path / "t.pdf")
    assert pdf.read_bytes()[:4] == b"%PDF"


def test_android_docedit_builtin(android, client):
    import docedit
    assert docedit.caps()["builtin"]
    r = client.post("/api/doc/save", data={"html": "<h1>Привет</h1><p>Текст</p>", "format": "docx", "name": "d"})
    assert r.status_code == 200, r.get_json()
    r = client.post("/api/doc/save", data={"html": "<p>x</p>", "format": "odt", "name": "d"})
    assert r.status_code == 400


def test_android_download_in_process(monkeypatch, tmp_path):
    import downloader
    from tasks import Task
    monkeypatch.setenv("UNIVERSAL_ANDROID", "1")
    calls = []

    def fake(urls, o, out, cookie, emit):
        emit({"progress": 0.5, "message": "Качаю"})
        (out / "v.mp4").write_bytes(b"x")
        emit({"done": True})
        calls.append(urls)
    monkeypatch.setattr(downloader, "_download_all", fake)
    monkeypatch.setattr(downloader, "yt_dlp", object())
    (tmp_path / "out").mkdir()
    t = Task("d", "dl", tmp_path, "dl")
    downloader.run_download_task(t, ["https://example.com/v"], {}, None)
    assert calls and t.done == 1 and (tmp_path / "out" / "v.mp4").exists()


def _https_server(tmp_path):
    """Локальный HTTPS с самоподписанным сертификатом: (адрес, путь к сертификату, остановка)."""
    import datetime
    import http.server
    import ssl
    import threading
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import ec
    from cryptography.x509.oid import NameOID
    import ipaddress
    key = ec.generate_private_key(ec.SECP256R1())
    name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "127.0.0.1")])
    now = datetime.datetime.now(datetime.timezone.utc)
    cert = (x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key())
            .serial_number(x509.random_serial_number()).not_valid_before(now - datetime.timedelta(days=1))
            .not_valid_after(now + datetime.timedelta(days=1))
            .add_extension(x509.SubjectAlternativeName([x509.IPAddress(ipaddress.ip_address("127.0.0.1"))]), False)
            .add_extension(x509.BasicConstraints(ca=True, path_length=None), True)
            .sign(key, hashes.SHA256()))
    crt, pem = tmp_path / "c.pem", tmp_path / "k.pem"
    crt.write_bytes(cert.public_bytes(serialization.Encoding.PEM))
    pem.write_bytes(key.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                                      serialization.NoEncryption()))

    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.end_headers()
            self.wfile.write(b"ok")

        def log_message(self, *a):
            pass

    srv = http.server.HTTPServer(("127.0.0.1", 0), H)
    sctx = ssl.SSLContext(ssl.PROTOCOL_TLS_SERVER)
    sctx.load_cert_chain(crt, pem)
    srv.socket = sctx.wrap_socket(srv.socket, server_side=True)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return f"https://127.0.0.1:{srv.server_port}/", crt, srv.shutdown


def test_net_falls_back_to_next_cert_check(tmp_path, monkeypatch):
    """Системная проверка не знает сертификат (старый Windows) -> следующий способ; проверка не отключается."""
    import ssl
    import net
    url, crt, stop = _https_server(tmp_path)
    try:
        trusting = ssl.create_default_context()
        trusting.load_verify_locations(crt)
        monkeypatch.setattr(net, "_contexts", [("system", ssl.create_default_context()), ("bundled", trusting)])
        monkeypatch.setattr(net, "_good", 0)
        with net.urlopen(url, timeout=10) as r:
            assert r.read() == b"ok"
        assert net._good == 1          # дальше сразу рабочим способом

        monkeypatch.setattr(net, "_contexts", [("system", ssl.create_default_context())])
        monkeypatch.setattr(net, "_good", 0)
        with pytest.raises(net.CertError) as e:
            net.urlopen(url, timeout=10)
        assert "сертификат" in net.reason(e.value) and "антивирус" in net.reason(e.value)
    finally:
        stop()


def test_engine_download_error_names_the_reason(tmp_path, monkeypatch):
    import urllib.error
    import engines
    import net

    def boom(*a, **k):
        raise urllib.error.URLError(ConnectionRefusedError(111, "Connection refused"))
    monkeypatch.setattr(net, "urlopen", boom)
    with pytest.raises(engines.EngineError) as e:
        engines._download_once("https://example.invalid/x.zip", tmp_path / "x.zip", None, 0, 1)
    assert "Connection refused" in str(e.value)


def test_engines_not_redownloaded_when_paths_json_broken(engine_env, monkeypatch):
    """paths.json испорчен (приложение закрыли во время записи): движки на диске находятся сами."""
    engines, srv_dir, base = engine_env
    with zipfile.ZipFile(srv_dir / "ff.zip", "w") as z:
        z.writestr("b/bin/ffmpeg.exe", "x")
        z.writestr("b/bin/ffprobe.exe", "x")
    monkeypatch.setitem(engines.CATALOG, "ffmpeg", {**engines.CATALOG["ffmpeg"], "url": base + "/ff.zip"})
    monkeypatch.setenv("PATH", "/usr/bin:/bin")
    first = engines.install("ffmpeg")
    assert not list((engines.engines_dir() / "downloads").glob("*")), "архив после установки не нужен"

    (engines.engines_dir() / "paths.json").write_text('{"ffmpeg": "C:\\\\', encoding="utf-8")
    (srv_dir / "ff.zip").unlink()                        # скачать заново уже нельзя
    monkeypatch.setenv("PATH", "/usr/bin:/bin")
    assert "ffmpeg" not in engines.missing()
    assert engines.install("ffmpeg") == first            # без загрузки
    assert json.loads((engines.engines_dir() / "paths.json").read_text(encoding="utf-8"))["ffmpeg"] == first


def test_tesseract_kept_when_languages_fail(engine_env, monkeypatch):
    """Словари не скачались: сам Tesseract остаётся, в следующий раз докачиваются только словари."""
    import time
    engines, srv_dir, base = engine_env
    monkeypatch.setattr(time, "sleep", lambda s: None)
    with zipfile.ZipFile(srv_dir / "t.zip", "w") as z:
        z.writestr("Tesseract-OCR/tesseract.exe", "x")
    monkeypatch.setitem(engines.CATALOG, "tesseract", {**engines.CATALOG["tesseract"], "url": base + "/t.zip",
                                                       "nsis": False, "langs": ["rus"]})
    monkeypatch.setattr(engines, "TESSDATA_URL", base + "/{lang}.traineddata")
    monkeypatch.setenv("PATH", "/usr/bin:/bin")
    with pytest.raises(engines.EngineError):
        engines.install("tesseract")
    assert engines._load_paths().get("tesseract")        # сам движок запомнен
    assert "tesseract" in engines.missing()              # но без словаря он не готов

    (srv_dir / "rus.traineddata").write_bytes(b"r" * 2000)
    (srv_dir / "t.zip").unlink()
    d = engines.install("tesseract")
    assert (Path(d) / "tessdata" / "rus.traineddata").stat().st_size == 2000
    assert "tesseract" not in engines.missing()


@pytest.fixture
def range_server():
    """HTTP с Range/If-Range и ETag, как у GitHub: (адрес, данные, список заголовков Range, смена версии)."""
    import http.server
    import threading
    state = {"data": bytes(range(256)) * 12000, "etag": '"v1"', "ranges": []}

    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            data, rng = state["data"], self.headers.get("Range")
            state["ranges"].append(rng)
            if rng and self.headers.get("If-Range") in (None, state["etag"]):
                start = int(rng.split("=")[1].rstrip("-"))
                if start >= len(data):
                    self.send_response(416)
                    self.end_headers()
                    return
                self.send_response(206)
                body = data[start:]
            else:
                self.send_response(200)
                body = data
            self.send_header("ETag", state["etag"])
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *a):
            pass

    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_port}/f.bin", state
    srv.shutdown()


def test_download_resumes_after_restart(tmp_path, range_server):
    import net
    url, state = range_server
    data = state["data"]
    dest = tmp_path / "f.bin"
    # прошлый запуск успел скачать треть и закрылся
    (tmp_path / "f.bin.part").write_bytes(data[:len(data) // 3])
    (tmp_path / "f.bin.part.json").write_text(json.dumps({"url": url, "etag": '"v1"'}), encoding="utf-8")
    seen = []
    net.download(url, dest, progress=lambda got, total: seen.append((got, total)))
    assert dest.read_bytes() == data
    assert state["ranges"] == [f"bytes={len(data) // 3}-"]          # докачан только остаток
    assert seen[-1] == (len(data), len(data))
    assert not (tmp_path / "f.bin.part").exists() and not (tmp_path / "f.bin.part.json").exists()


def test_download_restarts_when_file_changed(tmp_path, range_server):
    """На сервере новая версия (ffmpeg latest меняется каждый день): кусок старой не приклеиваем."""
    import net
    url, state = range_server
    (tmp_path / "f.bin.part").write_bytes(b"old" * 1000)
    (tmp_path / "f.bin.part.json").write_text(json.dumps({"url": url, "etag": '"v0"'}), encoding="utf-8")
    net.download(url, tmp_path / "f.bin")
    assert (tmp_path / "f.bin").read_bytes() == state["data"]


def test_download_keeps_part_when_cancelled(tmp_path, range_server):
    import net
    url, state = range_server

    class Stop(Exception):
        pass
    calls = []

    def cancel():
        calls.append(1)
        if len(calls) > 1:
            raise Stop
    with pytest.raises(Stop):
        net.download(url, tmp_path / "f.bin", check_cancel=cancel)
    part = tmp_path / "f.bin.part"
    size = part.stat().st_size
    assert 0 < size < len(state["data"])
    net.download(url, tmp_path / "f.bin")                # следующий запуск продолжает с того же места
    assert (tmp_path / "f.bin").read_bytes() == state["data"]
    assert state["ranges"][-1] == f"bytes={size}-"


def test_engine_install_survives_locked_archive(engine_env, monkeypatch):
    """Windows держит архив после распаковки (7-Zip, антивирус): установка не падает, архив уберётся потом."""
    import pathlib
    engines, srv_dir, base = engine_env
    with zipfile.ZipFile(srv_dir / "ff.zip", "w") as z:
        z.writestr("b/ffmpeg.exe", "x")
        z.writestr("b/ffprobe.exe", "x")
    monkeypatch.setitem(engines.CATALOG, "ffmpeg", {**engines.CATALOG["ffmpeg"], "url": base + "/ff.zip"})
    monkeypatch.setenv("PATH", "/usr/bin:/bin")
    real_unlink = pathlib.Path.unlink

    def locked(self, *a, **k):
        if self.parent.name == "downloads" and self.suffix == ".zip":
            raise PermissionError(32, "The process cannot access the file because it is being used by another process")
        return real_unlink(self, *a, **k)
    monkeypatch.setattr(pathlib.Path, "unlink", locked)
    d = engines.install("ffmpeg")
    assert "ffmpeg" not in engines.missing()
    leftover = list((engines.engines_dir() / "downloads").glob("ffmpeg-*"))
    assert leftover                                   # остался: файл «занят»

    monkeypatch.setattr(pathlib.Path, "unlink", real_unlink)
    assert engines.install("ffmpeg") == d             # повторно — без загрузки, остатки убраны
    assert not list((engines.engines_dir() / "downloads").glob("ffmpeg-*"))


# ---------------------------------------------------------------- «Скачать» → «Страницу сайта»

@pytest.fixture
def site(tmp_path):
    """Локальный сайт: страница со стилями, @import, шрифтом, картинками, скриптами; /secret — только с cookie."""
    import http.server
    import io
    import threading
    from PIL import Image
    root = tmp_path / "site"
    root.mkdir()

    def png(color):
        b = io.BytesIO()
        Image.new("RGB", (8, 8), color).save(b, "PNG")
        return b.getvalue()
    (root / "pic.png").write_bytes(png("red"))
    (root / "pic2.png").write_bytes(png("blue"))
    (root / "bg.png").write_bytes(png("green"))
    (root / "icon.png").write_bytes(png("black"))
    (root / "f.woff2").write_bytes(b"wOF2" + b"\x00" * 200)
    (root / "extra.css").write_text(".x { color: blue }", encoding="utf-8")
    (root / "style.css").write_text(
        '@import "extra.css";\nbody { background: url(bg.png); font-family: F }\n'
        "@font-face { font-family: F; src: url(f.woff2) format('woff2') }\nh1 > span { color: red }\n", encoding="utf-8")
    (root / "index.html").write_text("""<!doctype html><html><head><meta charset="utf-8"><title>Тестовая  страница</title>
<link rel="stylesheet" href="style.css"><link rel="icon" href="icon.png">
<style>.inl { background: url(bg.png) } p > b { color: red }</style>
<script src="app.js"></script></head>
<body onload="x()"><h1 style="background:url('bg.png')">Привет</h1><p>Текст <b>жирный</b></p>
<img src="pic.png" srcset="pic.png 1x, pic2.png 2x"><img src="missing.png">
<picture><source srcset="pic2.png"><img src="pic.png"></picture>
<a href="page2.html">дальше</a> <a href="javascript:alert(1)">js</a>
<noscript><p id="ns">без скриптов</p></noscript><video src="v.mp4"></video>
<script>alert(1)</script></body></html>""", encoding="utf-8")

    class H(http.server.SimpleHTTPRequestHandler):
        def __init__(self, *a, **k):
            super().__init__(*a, directory=str(root), **k)

        def do_GET(self):
            if self.path == "/secret":
                if "sid=1" not in (self.headers.get("Cookie") or ""):
                    self.send_response(403)
                    self.end_headers()
                    return
                body = "<html><head><title>Секрет</title></head><body>для своих</body></html>".encode()
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
                return
            if self.path == "/go":
                self.send_response(302)
                self.send_header("Location", "/secret")
                self.end_headers()
                return
            if self.path == "/gz":
                import gzip
                body = gzip.compress("<html><head><title>Сжатая</title></head><body>gzip</body></html>".encode())
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Encoding", "gzip")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
                return
            super().do_GET()

        def log_message(self, *a):
            pass

    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{srv.server_port}"
    srv.shutdown()


def test_page_builtin_inlines_everything(site, tmp_path):
    import webpage
    out = tmp_path / "out"
    out.mkdir()
    f = webpage._save_builtin(site + "/index.html", out, webpage.read_opts({}), None, None)
    assert f.name == "Тестовая страница.html"
    h = f.read_text(encoding="utf-8")
    assert "<script" not in h and "onload" not in h and "alert(1)" not in h and "javascript:" not in h
    assert "<noscript" not in h and 'id="ns"' in h                      # запасной вариант страницы виден
    assert "<video" not in h and "<source" not in h and "srcset" not in h
    assert 'rel="stylesheet"' not in h and "@import" not in h and "style.css" not in h
    assert ".x { color: blue }" in h                                    # @import вшит
    assert "p > b" in h and "h1 > span" in h                            # «>» в стилях не превратился в &gt;
    assert h.count("data:image/png;base64,") >= 4                        # картинки, фон, значок
    assert "data:font/woff2;base64," in h or "data:application/font-woff2;base64," in h
    assert "url(bg.png)" not in h and "url('bg.png')" not in h
    assert f'src="{site}/missing.png"' in h                              # не открылась: остался полный адрес
    assert f'href="{site}/page2.html"' in h
    assert '<meta charset="utf-8"/>' in h or '<meta charset="utf-8">' in h
    assert "Сохранено Universal" in h


def test_page_builtin_options(site, tmp_path):
    import webpage
    out = tmp_path / "out"
    out.mkdir()
    o = webpage.read_opts({"no_images": "1", "no_fonts": "1", "no_js": "0"})
    h = webpage._save_builtin(site + "/index.html", out, o, None, None).read_text(encoding="utf-8")
    assert "data:image/png;base64," not in h and "data:font" not in h
    assert webpage.PIXEL in h
    assert "alert(1)" in h and f'src="{site}/app.js"' in h               # скрипты оставлены, адрес полный


def test_page_builtin_cookies_redirect_gzip(site, tmp_path):
    import webpage
    out = tmp_path / "out"
    out.mkdir()
    port = site.rsplit(":", 1)[1]
    with pytest.raises(webpage.PageError, match="403"):
        webpage._save_builtin(site + "/go", out, webpage.read_opts({}), None, None)
    cookies = tmp_path / "cookies.txt"
    cookies.write_text("# Netscape HTTP Cookie File\n127.0.0.1\tFALSE\t/\tFALSE\t0\tsid\t1\n", encoding="utf-8")
    f = webpage._save_builtin(site + "/go", out, webpage.read_opts({}), cookies, None)   # через переадресацию
    assert f.name == "Секрет.html" and "для своих" in f.read_text(encoding="utf-8"), port
    g = webpage._save_builtin(site + "/gz", out, webpage.read_opts({}), None, None)
    assert g.name == "Сжатая.html" and "gzip" in g.read_text(encoding="utf-8")


def test_page_builtin_file_link_saved_as_is(site, tmp_path):
    import webpage
    out = tmp_path / "out"
    out.mkdir()
    f = webpage._save_builtin(site + "/pic.png", out, webpage.read_opts({}), None, None)
    assert f.suffix == ".png" and f.read_bytes()[:4] == b"\x89PNG"


def test_page_route_saves_html(site, client, monkeypatch):
    monkeypatch.delenv("UNIVERSAL_ANDROID", raising=False)
    r = client.post("/api/dl", data={"mode": "page", "urls": f"{site}/index.html\n{site}/gz", "builtin": "1"})
    d = wait_task(client, r.get_json()["task"])
    assert d["status"] == "done" and d["done"] == 2 and not d["failed"], d
    assert sorted(d["names"]) == ["Сжатая.html", "Тестовая страница.html"]
    body = client.get(f"/api/file/{d['job']}/Тестовая страница.html").data.decode("utf-8")
    assert "data:image/png;base64," in body and "<script" not in body
    bad = client.post("/api/dl", data={"mode": "page", "urls": "просто слова"})
    assert bad.status_code == 400
    r = client.post("/api/dl", data={"mode": "page", "urls": "http://127.0.0.1:9/none", "builtin": "1"})
    d = wait_task(client, r.get_json()["task"])
    assert d["status"] == "error" and "не открылась" in d["error"]


@pytest.mark.skipif(not shutil.which("monolith"), reason="monolith не установлен")
def test_page_with_real_monolith(site, tmp_path):
    import webpage
    out = tmp_path / "out"
    out.mkdir()
    f = webpage._save_monolith(shutil.which("monolith"), site + "/index.html", out, webpage.read_opts({}), None, None)
    assert f.name == "Тестовая страница.html"
    h = f.read_text(encoding="utf-8")
    assert "alert(1)" not in h and "onload" not in h and "app.js" not in h      # пустой <script></script> monolith оставляет
    assert "data:image/png;base64," in h and "style.css" not in h
    import base64
    import re
    css = " ".join(re.findall(r"<style[^>]*>(.*?)</style>", h, re.S))
    for payload in re.findall(r"data:text/css;base64,([A-Za-z0-9+/=]+)", h):      # monolith вшивает стили как data:-ссылку
        css += " " + base64.b64decode(payload).decode("utf-8", "replace")
    assert "bg.png" not in css and "data:image/png;base64," in css              # картинка из стилей вшита


def test_engine_single_file_install(engine_env, monkeypatch):
    """monolith — один исполняемый файл без архива: ложится в engines/monolith и находится сам."""
    engines, srv_dir, base = engine_env
    (srv_dir / "monolith-bin").write_bytes(b"\x7fELF" + b"\x00" * 200_000)
    monkeypatch.setitem(engines.CATALOG, "monolith", {**engines.CATALOG["monolith"], "url": base + "/monolith-bin",
                                                       "asset_fn": None})
    monkeypatch.setenv("PATH", "/usr/bin:/bin")
    d = engines.install("monolith")
    exe = Path(d) / engines._exe_names("monolith")[0]
    assert exe.is_file() and exe.stat().st_size > 100_000
    if sys.platform != "win32":
        assert os.access(exe, os.X_OK)
    assert "monolith" not in engines.missing()
    assert not list((engines.engines_dir() / "downloads").glob("monolith-*"))
    (srv_dir / "tiny").write_bytes(b"<html>404</html>")
    monkeypatch.setitem(engines.CATALOG, "monolith", {**engines.CATALOG["monolith"], "url": base + "/tiny"})
    engines.remove("monolith")
    with pytest.raises(engines.EngineError, match="не тот файл"):
        engines.install("monolith")


def test_fresh_ytdlp_not_activated_on_python_310(tmp_path, monkeypatch):
    """Свежий yt-dlp (3.11+) на Python 3.10 не подключаем и не предлагаем: он не импортируется, и всё скачивание умерло бы."""
    import engines
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path))
    z = tmp_path / "engines" / "yt-dlp.zip"
    z.parent.mkdir(parents=True)
    z.write_bytes(b"PK")
    (z.parent / "paths.json").write_text(json.dumps({"ytdlp": str(z)}), encoding="utf-8")
    monkeypatch.setattr(sys, "path", list(sys.path))
    monkeypatch.setattr(sys, "version_info", (3, 10, 12, "final", 0))
    engines.activate()
    assert str(z) not in sys.path and not engines.can_install("ytdlp")
    monkeypatch.setattr(sys, "version_info", (3, 12, 1, "final", 0))
    engines.activate()
    assert str(z) in sys.path and engines.can_install("ytdlp")


# ---------------------------------------------------------------- читалка (TextFold, вшит в third_party)

@pytest.fixture(scope="module")
def reader_env(tmp_path_factory):
    """Читалка запускается один раз на модуль: данные во временной папке, модели не ищем и не вызываем."""
    import reader
    if not reader.available():
        pytest.skip("читалке не хватает библиотек: " + ", ".join(reader.missing()))
    mp = pytest.MonkeyPatch()
    mp.setenv("TEXTFOLD_DATA", str(tmp_path_factory.mktemp("textfold")))
    mp.setenv("TEXTFOLD_NO_MODELS", "1")
    mp.setenv("UNIVERSAL_DESKTOP", "1")
    yield reader
    mp.undo()


def test_reader_imports_text_and_serves_page(reader_env):
    import urllib.request
    reader = reader_env
    url = reader.start()
    assert url.startswith("http://127.0.0.1:") and reader.start() == url          # второй запуск тот же
    r = reader.import_text("# Заголовок статьи\n\nПервый абзац про дело.\n\n## Раздел\n\nВторой абзац с числом 42.", "статья.md")
    assert r["paper"].startswith("md-") and r["title"] == "Заголовок статьи" and r["url"] == f"{url}/p/{r['paper']}"
    with urllib.request.urlopen(r["url"], timeout=10) as page:
        html = page.read().decode("utf-8")
    assert page.status == 200 and "Заголовок статьи" in html
    with pytest.raises(reader.ReaderError, match="Пустой"):
        reader.import_text("  \n ")


def test_reader_imports_word_and_rejects_unknown(reader_env, tmp_path):
    import docx
    reader = reader_env
    d = docx.Document()
    d.add_heading("Отчёт", 1)
    d.add_paragraph("Текст отчёта про важное дело.")
    f = tmp_path / "отчёт.docx"
    d.save(f)
    r = reader.import_file(f)
    assert r["paper"].startswith("doc-") and r["url"].endswith(r["paper"])
    bad = tmp_path / "x.mp3"
    bad.write_bytes(b"ID3")
    with pytest.raises(reader.ReaderError, match="не открывает"):
        reader.import_file(bad)
    scan = tmp_path / "скан.png"
    scan.write_bytes(b"\x89PNG\r\n\x1a\n" + b"0" * 50)
    with pytest.raises(reader.ReaderError, match="MinerU"):
        reader.import_file(scan)


def test_reader_routes(reader_env, client, monkeypatch):
    import io
    st = client.get("/api/reader").get_json()
    assert st["available"] and st["usable"] and st["url"].startswith("http://127.0.0.1:")
    eng = client.get("/api/formats").get_json()["engines"]["reader"]
    assert eng["usable"] is True and eng["local"] is True and eng["missing"] == []
    r = client.post("/api/reader/open", data={"text": "# Привет\n\nТекст для чтения.", "name": "n.md"})
    d = r.get_json()
    assert r.status_code == 200 and d["url"].endswith("/p/" + d["paper"])
    r = client.post("/api/reader/open", data={"files": (io.BytesIO("# Из файла\n\nАбзац.".encode()), "файл.md")},
                    content_type="multipart/form-data")
    assert r.status_code == 200 and r.get_json()["title"] == "Из файла"
    assert client.post("/api/reader/open", data={"text": " "}).status_code == 400
    # телефону через сервер читалку не показываем: 127.0.0.1 компьютера ему недоступен
    import base64
    import phoneserver
    monkeypatch.setattr(phoneserver, "password", lambda: "pw")
    auth = {"Authorization": "Basic " + base64.b64encode(b"x:pw").decode()}
    phone = {"environ_overrides": {phoneserver.MARK: True}, "headers": auth}
    st = client.get("/api/reader", **phone).get_json()
    assert st["usable"] is False and st["local"] is False
    assert client.post("/api/reader/open", data={"text": "x"}, **phone).status_code == 403


def test_reader_unavailable_on_android(monkeypatch):
    import reader
    monkeypatch.setenv("UNIVERSAL_ANDROID", "1")
    assert not reader.available() and not reader.reachable()


def test_reader_missing_libraries_hint(monkeypatch):
    import importlib.util
    import reader
    real = importlib.util.find_spec
    monkeypatch.setattr(importlib.util, "find_spec", lambda n, *a, **k: None if n in ("trafilatura", "lxml") else real(n, *a, **k))
    assert {"trafilatura", "lxml"} <= set(reader.missing()) and not reader.available()
    cmd = reader.pip_command()
    assert cmd.startswith("pip install --user ") and '"trafilatura>=2.0"' in cmd and " lxml" in cmd
    monkeypatch.setattr(reader, "_port", 0)
    with pytest.raises(reader.ReaderError, match="не хватает"):
        reader.start()


# ---------------------------------------------------------------- документы в Markdown (MinerU)

FAKE_KIT = '''#!/usr/bin/env python3
import base64, os, sys
a = sys.argv[1:]
if a[:2] == ["models", "download"]:
    os.makedirs(os.path.join(os.environ["MINERU_HOME"], "models"), exist_ok=True)
    print("Fetching 2 files: 100%|██████████| 2/2", flush=True)
    sys.exit(0)
assert a[0] == "parse", a
src, out = a[1], a[a.index("-o") + 1]
tier = a[a.index("--tier") + 1]
if src.endswith("broken.pdf"):
    print("boom", file=sys.stderr); sys.exit(3)
stem = os.path.splitext(os.path.basename(src))[0]
png = base64.b64encode(b"\\x89PNG\\r\\n\\x1a\\n" + b"0" * 20).decode()
text = "" if src.endswith("empty.pdf") else f"# Заголовок {stem}\\n\\nТекст ({tier}).\\n\\n![](data:image/png;base64,{png})\\n\\n![рис](images/a.png)\\n"
os.makedirs(os.path.join(out, "images"), exist_ok=True)
open(os.path.join(out, "images", "a.png"), "wb").write(b"\\x89PNG\\r\\n\\x1a\\n" + b"1" * 20)
open(os.path.join(out, stem + ".md"), "w", encoding="utf-8").write(text)
'''

FAKE_UV = '''#!/usr/bin/env python3
import os, sys
a = sys.argv[1:]
log = os.environ.get("FAKE_UV_LOG")
if log:
    open(log, "a").write(" ".join(a) + "\\n")
if a[0] == "venv":
    os.makedirs(os.path.join(a[1], "bin"), exist_ok=True)
    open(os.path.join(a[1], "bin", "python"), "w").write("")
elif a[:2] == ["pip", "install"]:
    py = a[a.index("--python") + 1]
    kit = os.path.join(os.path.dirname(py), "mineru-kit")
    open(kit, "w").write(open(os.environ["FAKE_KIT_SRC"]).read())
    os.chmod(kit, 0o755)
    print("Installing packages:  50%|█████     | 1/2", flush=True)
'''


@pytest.fixture
def md_engine(tmp_path, monkeypatch):
    """Поддельные uv и mineru-kit: проверяем установку и конвертацию без гигабайтов (только Linux и macOS)."""
    if sys.platform == "win32":
        pytest.skip("поддельные исполняемые файлы только для Linux и macOS")
    import mdconv
    import reader
    reader.prepare()
    from textfold import engine
    home = tmp_path / "engine"
    monkeypatch.setattr(engine, "home", lambda: home)
    monkeypatch.setattr(engine, "legacy", lambda: tmp_path / "no-such-engine")
    uv = tmp_path / "uv"
    uv.write_text(FAKE_UV)
    uv.chmod(0o755)
    kit = tmp_path / "kit-src.py"
    kit.write_text(FAKE_KIT)
    monkeypatch.setattr(engine, "_uv", lambda e, log, stop: str(uv))
    monkeypatch.setenv("FAKE_KIT_SRC", str(kit))
    monkeypatch.setenv("FAKE_UV_LOG", str(tmp_path / "uv.log"))
    monkeypatch.delenv("UNIVERSAL_ANDROID", raising=False)
    monkeypatch.delenv("UNIVERSAL_DOCKER", raising=False)
    return mdconv, engine, home


def _task(tmp_path, kind="md"):
    from tasks import Task
    (tmp_path / "job" / "out").mkdir(parents=True, exist_ok=True)
    return Task("j1", kind, tmp_path / "job", "md")


def test_md_install_levels_and_status(md_engine, tmp_path):
    mdconv, engine, home = md_engine
    assert not mdconv.status()["basic"] and not mdconv.status()["full"]
    t = _task(tmp_path, "mdinstall")
    mdconv.install("basic", t)
    st = mdconv.status()
    assert st["basic"] and not st["full"] and t.done == 1 and t.progress == 1
    log = (tmp_path / "uv.log").read_text()
    assert "venv" in log and "--python-preference only-managed" in log and f"mineru=={engine.VERSION}" in log
    mdconv.install("full", _task(tmp_path, "mdinstall"))                     # модели докачиваются поверх базового
    st = mdconv.status()
    assert st["basic"] and st["full"] and (home / "mineru-home" / "models").is_dir()
    assert mdconv.uninstall()["basic"] is False and not home.exists()
    with pytest.raises(mdconv.MdError, match="Неизвестный"):
        mdconv.install("huge", _task(tmp_path))


def test_md_install_cancel(md_engine, tmp_path):
    from tasks import Cancelled
    mdconv, engine, home = md_engine
    t = _task(tmp_path, "mdinstall")
    t.cancel_event.set()
    with pytest.raises(Cancelled):
        mdconv.install("basic", t)


def test_md_convert_images_modes_and_errors(md_engine, tmp_path):
    mdconv, engine, home = md_engine
    mdconv.install("basic", _task(tmp_path, "mdinstall"))
    src = tmp_path / "in"
    src.mkdir()
    for n in ("doc.pdf", "empty.pdf", "broken.pdf"):
        (src / n).write_bytes(b"%PDF-1.4")

    t = _task(tmp_path)
    mdconv.run_md_task(t, [(src / "doc.pdf", "doc"), (src / "empty.pdf", "empty"), (src / "broken.pdf", "broken")],
                       mdconv.read_opts({"tier": "fast"}))
    out = tmp_path / "job" / "out"
    assert t.done == 1 and {f["name"] for f in t.failed} == {"empty.pdf", "broken.pdf"}
    md = (out / "doc.md").read_text(encoding="utf-8")
    assert "Заголовок doc" in md and "(flash)" in md        # быстрый режим = уровень flash у MinerU
    assert "base64" not in md                                         # вшитая картинка вынесена в файл
    assert (out / "doc_images" / "image1.png").read_bytes()[:4] == b"\x89PNG"
    assert "doc_images/a.png" in md and (out / "doc_images" / "a.png").is_file()   # и та, что лежала файлом
    assert "попробуй точный режим" in next(f["error"] for f in t.failed if f["name"] == "empty.pdf")
    assert "MinerU не справился" in next(f["error"] for f in t.failed if f["name"] == "broken.pdf")

    for mode, check in (("inline", lambda m, o: "base64" in m),
                        ("none", lambda m, o: "![" not in m and not (o / "doc_images").exists())):
        t = _task(tmp_path / mode)
        mdconv.run_md_task(t, [(src / "doc.pdf", "doc")], mdconv.read_opts({"tier": "fast", "images": mode}))
        o = tmp_path / mode / "job" / "out"
        assert check((o / "doc.md").read_text(encoding="utf-8"), o), mode


def test_md_accurate_needs_models(md_engine, tmp_path):
    mdconv, engine, home = md_engine
    src = tmp_path / "s.pdf"
    src.write_bytes(b"%PDF")
    with pytest.raises(mdconv.MdError, match="Точному режиму нужны модели"):
        mdconv.run_md_task(_task(tmp_path), [(src, "s")], mdconv.read_opts({"tier": "accurate"}))
    mdconv.install("full", _task(tmp_path, "mdinstall"))
    t = _task(tmp_path / "x")
    mdconv.run_md_task(t, [(src, "s")], mdconv.read_opts({"tier": "accurate"}))
    assert "(advanced)" in (tmp_path / "x" / "job" / "out" / "s.md").read_text(encoding="utf-8")   # точный = advanced
    with pytest.raises(mdconv.MdError, match="не установлен"):
        mdconv.uninstall()
        mdconv.run_md_task(_task(tmp_path / "y"), [(src, "s")], mdconv.read_opts({"tier": "fast"}))


def test_md_routes(md_engine, client, tmp_path, monkeypatch):
    import io
    import phoneserver
    mdconv, engine, home = md_engine
    st = client.get("/api/md/engine").get_json()
    assert st["supported"] and not st["basic"] and st["local"]
    assert client.get("/api/formats").get_json()["engines"]["mineru"]["supported"] is True
    # установка: задача, которую можно опросить, как и остальные
    r = client.post("/api/md/install", data={"level": "basic"})
    d = wait_task(client, r.get_json()["task"])
    assert d["status"] == "done", d
    assert client.get("/api/md/engine").get_json()["basic"] is True
    assert client.post("/api/md/install", data={"level": "huge"}).status_code == 400
    # конвертация: нужные файлы приходят в одном архиве
    r = client.post("/api/md", data={"tier": "fast", "files": (io.BytesIO(b"%PDF-1.4"), "отчёт.pdf")}, content_type="multipart/form-data")
    d = wait_task(client, r.get_json()["task"])
    assert d["status"] == "done" and d["done"] == 1, d
    assert "отчёт.md" in d["names"]
    assert "Заголовок" in client.get(f"/api/file/{d['job']}/отчёт.md").data.decode("utf-8")
    assert client.post("/api/md", data={"files": (io.BytesIO(b"x"), "музыка.mp3")}, content_type="multipart/form-data").status_code == 400
    assert client.post("/api/md", data={}).status_code == 400
    # телефону через сервер установка недоступна, а нет ли уже поставленного движка — не его дело
    monkeypatch.setattr(phoneserver, "password", lambda: "pw")
    import base64
    phone = {"environ_overrides": {phoneserver.MARK: True}, "headers": {"Authorization": "Basic " + base64.b64encode(b"x:pw").decode()}}
    assert client.post("/api/md/install", data={"level": "basic"}, **phone).status_code == 403
    assert client.get("/api/md/engine", **phone).get_json()["local"] is False
    assert client.post("/api/md/uninstall", **phone).status_code == 403


def test_md_unsupported_on_android(monkeypatch, client):
    import mdconv
    monkeypatch.setenv("UNIVERSAL_ANDROID", "1")
    assert not mdconv.supported() and mdconv.status()["supported"] is False
    assert client.post("/api/md/install", data={"level": "basic"}).status_code == 400


# ---------------------------------------------------------------- голос в текст: разделение по голосам

def test_diarize_assign_and_join():
    import diarize
    segs = [(0.0, 2.0, "Привет, как дела?"), (2.2, 4.0, "Нормально."), (4.1, 5.0, "А у тебя?"), (6.5, 7.0, "Тоже."), (9.0, 9.5, "Пока.")]
    found = [(0.0, 2.1, 7), (2.1, 4.05, 3), (4.05, 6.0, 7), (6.0, 7.2, 3)]       # номера у модели произвольные
    who = diarize.assign(segs, found)
    assert who == [0, 1, 0, 1, 1]                      # номера по порядку появления; «Пока» в паузе -> ближайший говорящий (3)
    assert diarize.join(segs, who) == ("Спикер 1: Привет, как дела?\n\nСпикер 2: Нормально.\n\nСпикер 1: А у тебя?\n\n"
                                       "Спикер 2: Тоже. Пока.")
    assert diarize.assign(segs, []) == [0] * 5
    assert diarize.read_speakers("auto") == -1 and diarize.read_speakers("3") == 3 and diarize.read_speakers("") == 0
    assert diarize.read_speakers("0") == 0 and diarize.read_speakers("99") == -1 and diarize.read_speakers("abc") == -1


def _stt_job(tmp_path, monkeypatch, who_segments, found):
    import diarize
    import stt
    from tasks import Task
    monkeypatch.setattr(stt, "available", lambda: True)
    monkeypatch.setattr(diarize, "available", lambda: True)
    monkeypatch.setattr(stt, "ensure_model", lambda key, task=None: tmp_path)
    monkeypatch.setattr(stt, "_model", lambda key: None)
    monkeypatch.setattr(stt, "load_audio", lambda src: b"audio")
    calls = []
    monkeypatch.setattr(stt, "transcribe", lambda src, key, lang, task=None, base=0, span=1, audio=None:
                        (calls.append(audio) or who_segments, "ru", 9.0))
    monkeypatch.setattr(diarize, "ensure_models", lambda task=None, base=0, span=0: tmp_path)
    monkeypatch.setattr(diarize, "turns", lambda audio, n=-1: found if not isinstance(found, Exception) else (_ for _ in ()).throw(found))
    job = tmp_path / "job"
    (job / "out").mkdir(parents=True, exist_ok=True)
    src = tmp_path / "meet.ogg"
    src.write_bytes(b"x")
    task = Task("m1", "stt", job, "text")
    return stt, task, src, job, calls


def test_stt_with_speakers_labels_text_and_srt(tmp_path, monkeypatch):
    segs = [(0.0, 2.0, "Добрый день."), (2.2, 4.0, "Здравствуйте."), (4.2, 6.0, "Начнём.")]
    stt, task, src, job, calls = _stt_job(tmp_path, monkeypatch, segs, [(0.0, 2.1, 0), (2.1, 4.1, 1), (4.1, 6.0, 0)])
    stt.run_stt_task(task, [(src, "meet")], stt.read_opts({"model": "fast", "srt": "1", "speakers": "auto"}))
    txt = (job / "out" / "meet.txt").read_text(encoding="utf-8")
    assert txt.strip() == "Спикер 1: Добрый день.\n\nСпикер 2: Здравствуйте.\n\nСпикер 1: Начнём."
    srt = (job / "out" / "meet.srt").read_text(encoding="utf-8")
    assert "[Спикер 2] Здравствуйте." in srt and srt.count("[Спикер") == 3
    r = task.result["transcripts"][0]
    assert r["speakers"] == 2 and r["text"].startswith("Спикер 1:")
    assert calls == [b"audio"]                          # звук прочитан один раз и передан в распознавание


def test_stt_single_speaker_has_no_labels_and_failure_keeps_text(tmp_path, monkeypatch):
    import diarize
    segs = [(0.0, 2.0, "Только я."), (2.5, 4.0, "Один говорю.")]
    stt, task, src, job, _ = _stt_job(tmp_path, monkeypatch, segs, [(0.0, 4.0, 5)])
    stt.run_stt_task(task, [(src, "meet")], stt.read_opts({"model": "fast", "speakers": "2"}))
    assert (job / "out" / "meet.txt").read_text(encoding="utf-8").strip() == "Только я. Один говорю."
    assert task.result["transcripts"][0]["speakers"] == 1
    stt, task, src, job, _ = _stt_job(tmp_path / "b", monkeypatch, segs, diarize.DiarizeError("модель сломана"))
    stt.run_stt_task(task, [(src, "meet")], stt.read_opts({"model": "fast", "speakers": "auto"}))
    assert (job / "out" / "meet.txt").read_text(encoding="utf-8").strip() == "Только я. Один говорю."   # текст не потерян
    assert task.done == 1 and "Голоса разделить не удалось: модель сломана" in task.failed[0]["error"]


def test_stt_speakers_needs_sherpa(monkeypatch):
    import diarize
    import stt
    monkeypatch.setattr(stt, "available", lambda: True)
    monkeypatch.setattr(diarize, "available", lambda: False)
    with pytest.raises(stt.SttError, match="sherpa-onnx"):
        stt.read_opts({"model": "fast", "speakers": "auto"})
    assert stt.read_opts({"model": "fast"})["speakers"] == 0          # без галочки ничего не требуется
    assert stt.caps()["diarize"]["available"] is False and "sherpa-onnx" in stt.caps()["diarize"]["pip"]


@pytest.mark.skipif(not os.environ.get("DIARIZE_REAL"), reason="настоящие модели голосов: DIARIZE_REAL=1 (качает ≈ 30 МБ)")
def test_diarize_real_four_speakers(tmp_path, monkeypatch):
    """Запись, где говорят четверо: настоящая сегментация и отпечатки голосов находят четверых."""
    import wave
    import numpy as np
    import diarize
    if not diarize.available():
        pytest.skip("нет sherpa-onnx")
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path / "home"))
    diarize.ensure_models()
    assert diarize.downloaded()
    wav = tmp_path / "four.wav"
    import net
    net.download("https://github.com/k2-fsa/sherpa-onnx/releases/download/speaker-segmentation-models/0-four-speakers-zh.wav", wav)
    w = wave.open(str(wav))
    x = np.frombuffer(w.readframes(w.getnframes()), np.int16).astype(np.float32) / 32768
    found = diarize.turns(x, -1)
    assert len({s for _, _, s in found}) == 4
    assert len({s for _, _, s in diarize.turns(x, 2)}) == 2          # число говорящих можно задать


def _fake_bundle(base):
    """Офлайн-набор в миниатюре: один движок, одна модель Whisper, модели голосов."""
    b = base / "offline"
    (b / "engines" / "pandoc" / "x").mkdir(parents=True)
    (b / "engines" / "pandoc" / "x" / "pandoc.exe").write_text("exe")
    (b / "whisper" / "fast").mkdir(parents=True)
    (b / "whisper" / "fast" / "model.bin").write_text("m")
    (b / "whisper" / "fast" / "config.json").write_text("{}")
    (b / "diarization").mkdir()
    (b / "diarization" / "segmentation.int8.onnx").write_text("s")
    (b / "diarization" / "embedding.onnx").write_text("e")
    return b


def test_offline_seed_copies_when_portable(tmp_path, monkeypatch):
    import offline
    import stt
    import diarize
    b = _fake_bundle(tmp_path / "app")
    monkeypatch.setattr(offline, "bundle_dir", lambda: b)
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path / "home"))
    assert sorted(offline.seed()) == ["diarization", "engines/pandoc", "whisper/fast"]
    assert stt.downloaded("fast") and diarize.downloaded()
    assert (tmp_path / "home" / "engines" / "pandoc" / "x" / "pandoc.exe").is_file()
    assert (b / "whisper" / "fast" / "model.bin").is_file()        # портативная папка остаётся целой
    assert offline.seed() == []                                    # второй запуск ничего не делает


def test_offline_seed_moves_when_installed_and_keeps_user_files(tmp_path, monkeypatch):
    import offline
    app = tmp_path / "app"
    b = _fake_bundle(app)
    (app / "unins000.exe").write_text("u")                         # поставил установщик
    home = tmp_path / "home"
    (home / "whisper" / "fast").mkdir(parents=True)
    (home / "whisper" / "fast" / "mine.txt").write_text("мой", encoding="utf-8")     # у человека уже есть своя модель
    monkeypatch.setattr(offline, "bundle_dir", lambda: b)
    monkeypatch.setenv("UNIVERSAL_HOME", str(home))
    done = offline.seed()
    assert "whisper/fast" not in done and (home / "whisper" / "fast" / "mine.txt").is_file()
    assert not (home / "whisper" / "fast" / "model.bin").exists()
    assert (home / "engines" / "pandoc").is_dir() and not (b / "engines" / "pandoc").exists()   # перенесено, не скопировано


def test_offline_absent_is_noop(tmp_path, monkeypatch):
    import offline
    monkeypatch.setattr(offline, "bundle_dir", lambda: None)
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path / "home"))
    assert offline.seed() == [] and offline.pending() == []


# ---------- проверка на вирусы (ClamAV) ----------

EICAR = rb"X5O!P%@AP[4\PZX54(P^)7CC)7}$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*"
needs_clam = pytest.mark.skipif(not (shutil.which("clamscan") and shutil.which("freshclam")), reason="ClamAV не установлен")


def _home_with_test_db(tmp_path, monkeypatch):
    """Базы вирусов из одной сигнатуры (хеш EICAR): настоящие весят 300 МБ и качаются из сети."""
    import hashlib
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path / "home"))
    import av
    d = av.db_dir()
    d.mkdir(parents=True)
    (d / "test.hdb").write_text(f"{hashlib.md5(EICAR).hexdigest()}:{len(EICAR)}:Eicar-Test-Signature\n")
    return av


@needs_clam
def test_av_scan_finds_virus_and_keeps_clean(tmp_path, monkeypatch):
    av = _home_with_test_db(tmp_path, monkeypatch)
    job = tmp_path / "job"
    (job / "in").mkdir(parents=True)
    (job / "out").mkdir()
    a, b = job / "in" / "письмо.txt", job / "in" / "вирус.com"
    a.write_text("привет")
    b.write_bytes(EICAR)
    from tasks import Task
    t = Task("x", "av", job, "av")
    av.run_av_task(t, [(a, "письмо"), (b, "вирус")], av.read_opts({}))
    r = t.result["av"]
    assert r["scanned"] == 2 and r["clean"] == 1
    assert [i["virus"].split(".")[0] for i in r["infected"]] == ["Eicar-Test-Signature"]
    assert r["infected"][0]["name"] == "вирус.com"            # показываем своё имя, а не служебное f0002
    assert b.read_bytes() == EICAR and a.exists()              # ничего не удаляем и не меняем
    assert not (job / "scan").exists()


@needs_clam
def test_av_route_end_to_end(tmp_path, monkeypatch):
    av = _home_with_test_db(tmp_path, monkeypatch)
    import app as webapp
    c = webapp.app.test_client()
    st = c.get("/api/av/engine").get_json()
    assert st["available"] and st["db"]["ready"] and st["supported"]
    assert c.post("/api/av", data={}).status_code == 400
    r = c.post("/api/av", data={"files": [(io.BytesIO(EICAR), "eicar.txt"), (io.BytesIO(b"hello"), "ok.txt")]},
               content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    d = wait_task(c, r.get_json()["task"])
    assert d["status"] == "done", d
    assert d["result"]["av"]["infected"][0]["name"] == "eicar.txt" and d["result"]["av"]["clean"] == 1


def test_av_needs_engine_and_db(tmp_path, monkeypatch):
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path / "home"))
    import av
    import app as webapp
    monkeypatch.setattr(av, "clamscan", lambda: None)
    monkeypatch.setattr(av, "SYSTEM_DBS", ())
    c = webapp.app.test_client()
    r = c.post("/api/av", data={"files": [(io.BytesIO(b"x"), "a.txt")]}, content_type="multipart/form-data")
    assert r.status_code == 400 and "ClamAV" in r.get_json()["error"]
    assert c.get("/api/av/engine").get_json()["db"]["ready"] is False
    assert c.get("/api/formats").get_json()["engines"]["av"]["available"] is False


@pytest.mark.skipif(sys.platform == "win32", reason="поддельный freshclam — shell-скрипт")
def test_av_update_db_with_fake_freshclam(tmp_path, monkeypatch):
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path / "home"))
    import av
    from tasks import Task
    fake = tmp_path / "freshclam"
    fake.write_text('#!/bin/sh\nfor a in "$@"; do case $a in --datadir=*) d=${a#--datadir=};; --config-file=*) c=${a#--config-file=};; esac; done\n'
                    'grep -q "^DatabaseMirror database.clamav.net" "$c" || exit 3\n'
                    'echo "Downloading daily.cvd [100%]"; echo "x" > "$d/daily.cvd"\n')
    fake.chmod(0o755)
    monkeypatch.setattr(av, "freshclam", lambda: str(fake))
    job = tmp_path / "j"
    job.mkdir()
    t = Task("y", "avinstall", job, "av")
    av.update_db(t)
    assert av.db_info()["ready"] and (av.db_dir() / "daily.cvd").exists()
    bad = tmp_path / "freshclam2"
    bad.write_text("#!/bin/sh\necho 'ERROR: Can not reach database.clamav.net'\nexit 62\n")
    bad.chmod(0o755)
    (av.db_dir() / "daily.cvd").unlink()
    monkeypatch.setattr(av, "freshclam", lambda: str(bad))
    with pytest.raises(av.AvError) as e:
        av.update_db(t)
    assert "Не получилось скачать базы" in str(e.value) and "database.clamav.net" in str(e.value)


@needs_clam
@pytest.mark.skipif(not os.environ.get("AV_REAL"), reason="настоящие базы вирусов: AV_REAL=1 (качает ≈ 300 МБ)")
def test_av_real_database_and_eicar(tmp_path, monkeypatch):
    """freshclam скачивает настоящие базы нашей командой, и по ним находится тестовый вирус EICAR."""
    monkeypatch.setenv("UNIVERSAL_HOME", str(tmp_path / "home"))
    import av
    from tasks import Task
    job = tmp_path / "job"
    (job / "in").mkdir(parents=True)
    (job / "out").mkdir()
    t0 = Task("u", "avinstall", job, "av")
    av.update_db(t0)
    info = av.db_info()
    assert info["ready"] and info["mb"] > 100, info
    a, b = job / "in" / "a.txt", job / "in" / "eicar.com"
    a.write_text("hello")
    b.write_bytes(EICAR)
    t = Task("v", "av", job, "av")
    av.run_av_task(t, [(a, "a"), (b, "eicar")], av.read_opts({}))
    r = t.result["av"]
    assert r["clean"] == 1 and len(r["infected"]) == 1 and "EICAR" in r["infected"][0]["virus"].upper()


# ---------- базы данных: просмотр SQLite и DBeaver ----------

def _make_db(path):
    import sqlite3
    con = sqlite3.connect(path)
    con.execute('create table "люди и ""кавычки""" (id integer primary key, имя text, фото blob)')
    con.executemany('insert into "люди и ""кавычки""" (имя, фото) values (?, ?)', [("Аня", b"\x00\x01\x02"), ("Боря", None)] * 30)
    con.execute("create table empty (a)")
    con.execute('create view v as select имя from "люди и ""кавычки"""')
    con.commit()
    con.close()


def test_db_inspect_and_rows(tmp_path):
    import dbtools
    f = tmp_path / "t.db"
    _make_db(f)
    info = dbtools.inspect(f)
    names = {t["name"]: t for t in info["tables"]}
    assert names['люди и "кавычки"']["rows"] == 60 and names['люди и "кавычки"']["columns"] == ["id", "имя", "фото"]
    assert names["empty"]["rows"] == 0 and names["v"]["view"] and names["v"]["rows"] is None
    r = dbtools.rows(f, 'люди и "кавычки"')
    assert len(r["rows"]) == dbtools.PREVIEW_ROWS and r["total"] == 60
    assert r["rows"][0][1] == "Аня" and r["rows"][0][2] == "<3 байт>" and r["rows"][1][2] is None
    with pytest.raises(dbtools.DbError):
        dbtools.rows(f, "нет такой")
    junk = tmp_path / "x.db"
    junk.write_bytes(b"not a database at all")
    with pytest.raises(dbtools.DbError):
        dbtools.inspect(junk)
    # смотрим только на чтение: файл не меняется
    before = f.read_bytes()
    dbtools.inspect(f); dbtools.rows(f, "empty")
    assert f.read_bytes() == before


def test_db_routes_and_open_in_dbeaver(tmp_path, monkeypatch):
    import app as webapp
    import dbtools
    src = tmp_path / "my data.sqlite3"
    _make_db(src)
    c = webapp.app.test_client()
    assert c.post("/api/db", data={"files": [(io.BytesIO(b"x"), "a.txt")]}, content_type="multipart/form-data").status_code == 400
    r = c.post("/api/db", data={"files": [(io.BytesIO(src.read_bytes()), "my data.sqlite3"), (io.BytesIO("плохой файл".encode()), "bad.db")]},
               content_type="multipart/form-data")
    d = r.get_json()
    assert r.status_code == 200 and d["files"][0]["tables"] and "error" in d["files"][1]
    job, name = d["job"], d["files"][0]["name"]
    assert name.endswith(".sqlite3")                                  # расширение .sqlite3 не превращается в .bin
    rows = c.get("/api/db/rows", query_string={"job": job, "file": name, "table": "empty"}).get_json()
    assert rows["total"] == 0 and rows["columns"] == ["a"]
    assert c.get("/api/db/rows", query_string={"job": job, "file": "../x", "table": "t"}).status_code == 404
    assert c.get(f"/api/file/{job}/{name}").status_code == 200        # копию можно скачать
    # DBeaver нет: понятная ошибка
    monkeypatch.setattr(dbtools, "dbeaver_cmd", lambda: None)
    assert "не установлен" in c.post("/api/db/open", data={"job": job, "file": name}).get_json()["error"]
    if sys.platform != "win32":
        log = tmp_path / "args.txt"
        fake = tmp_path / "dbeaver"
        fake.write_text(f'#!/bin/sh\nfor a in "$@"; do echo "$a"; done > "{log}"\n')
        fake.chmod(0o755)
        monkeypatch.setattr(dbtools, "dbeaver_cmd", lambda: [str(fake)])
        assert c.post("/api/db/open", data={"job": job, "file": name}).get_json() == {"ok": True}
        for _ in range(50):
            if log.exists() and log.read_text():
                break
            time.sleep(0.1)
        args = log.read_text().splitlines()
        assert args[0] == "-con" and "driver=sqlite_jdbc" in args[1] and f"database={webapp.WORK_ROOT.resolve()}" in args[1] and name in args[1]
        assert "connect=true" in args[1]


def test_db_dbeaver_engine_entry():
    import engines
    spec = engines.CATALOG["dbeaver"]
    assert spec["exe"] == ["dbeaver.exe"] and "dbeaver" in spec["repo"]
    import re
    assert re.search(spec["asset"], "dbeaver-ce-24.3.0-win32.win32.x86_64.zip")
    assert not re.search(spec["asset"], "dbeaver-ce-24.3.0-linux.gtk.x86_64.tar.gz")
    assert "dbeaver" not in engines.AUTO and "clamav" not in engines.AUTO


def test_db_connection_arg_matches_dbeaver_spec():
    """Формат -con сверен с исходниками DBeaver: драйвер sqlite_jdbc, параметры через |, обратный слэш удваивается."""
    import dbtools
    from pathlib import PureWindowsPath
    arg = dbtools.connection_arg(PureWindowsPath(r"C:\Users\Аня\jobs\x.db"), "x.db")
    assert arg == r"name=Universal: x.db|driver=sqlite_jdbc|database=C:\\Users\\Аня\\jobs\\x.db|connect=true|openConsole=false"
    with pytest.raises(dbtools.DbError):
        dbtools.connection_arg("/tmp/a|b.db", "a")
